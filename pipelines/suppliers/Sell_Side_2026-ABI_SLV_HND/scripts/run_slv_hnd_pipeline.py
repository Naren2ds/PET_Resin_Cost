"""
Pipeline: Sell Side 2026 - ABI (El Salvador and Honduras)
Source: Sell_Side_2026-ABI May 2026 AMCOR (1).xlsx
Sheet: ICIS Price History
Data range: I13:R (row 14 to last row with data)
Destination: El Salvador and Honduras
Supplier: Amcor

TLC Formula:
  Total precio ABI = Index (N-4) + Adder + Logistics + LANDED + Finance fee

Mapping (from Mapping_Columns.xlsx / Sell_Side_2026-ABI Mar 2026 sheet):
  - PET Bottle Grade FOB China (Mid) N-4 -> Resin Index vPET
  - SC -> Tax
  - Adder -> Others
  - Logistics -> Freight
  - LANDED -> Tax
  - Finance fee V (SOFR +3,8%@185) -> Insurance
  - CME TERM SOFR (%) 180 -> Insurance (not in TLC calc)
  - Surcharge -> Others
  - Total precio ABI -> Total Landing Cost
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Path constants
# ---------------------------------------------------------------------------
PIPELINE_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[4]

SOURCE_FILE = REPO_ROOT / "data" / "raw" / "suppliers" / "Sell_Side_2026-ABI May 2026 AMCOR (1).xlsx"
INDEX_REFERENCE_CSV = REPO_ROOT / "data" / "processed" / "current" / "indexes" / "suppliers" / "icis_resin_index_reference_table.csv"
MAPPING_FILE = REPO_ROOT / "data" / "reference" / "mappings" / "Mapping_Columns.xlsx"

ARTIFACTS_DIR = PIPELINE_DIR / "artifacts"
EXTRACTION_DIR = ARTIFACTS_DIR / "extraction"
FRONT_END_DIR = ARTIFACTS_DIR / "front_end"
VALIDATION_DIR = ARTIFACTS_DIR / "validation"
FORECAST_DIR = ARTIFACTS_DIR / "forecast"

# ---------------------------------------------------------------------------
# Pipeline metadata
# ---------------------------------------------------------------------------
PIPELINE_NAME = "Sell_Side_2026-ABI_SLV_HND"
SOURCE_SHEET = "ICIS Price History"
HEADER_ROW = 13
DATA_START_ROW = 14

DEFAULT_METADATA = {
    "supplier": "Amcor",
    "destination_country": "El Salvador and Honduras",
    "sourcing_country": "China",
    "resin_index_type": "PET Bottle Grade FOB China Mid (N-4) USD/ton",
    "tlc_formula": "Total precio ABI = Index(N-4) + Adder + Logistics + LANDED + Finance fee",
}

# Column positions in the May 2026 workbook (1-indexed)
COL_DATE = 2        # B: Date
COL_INDEX = 9       # I: PET Bottle Grade FOB China (Mid) N-4
COL_SC = 10         # J: SC
COL_ADDER = 11      # K: Adder
COL_LOGISTICS = 12  # L: Logistics
COL_LANDED = 13     # M: LANDED
COL_FINANCE = 14    # N: Finance fee
COL_SOFR = 15       # O: CME TERM SOFR (%)
COL_SURCHARGE = 16  # P: Surcharge
COL_TOTAL = 18      # R: Total precio ABI

# Mapping: column index -> (raw_label, mapping_column, required_for_calc)
COMPONENT_COLUMNS = {
    COL_INDEX: ("PET Bottle Grade FOB China Mid (N-4)", "Resin Index vPET", True),
    COL_SC: ("SC", "Tax", True),
    COL_ADDER: ("Adder", "Others", True),
    COL_LOGISTICS: ("Logistics", "Freight ", True),
    COL_LANDED: ("LANDED", "Tax", True),
    COL_FINANCE: ("Finance fee V (SOFR +3,8%@185)", "Insurance", True),
    COL_SOFR: ("CME TERM SOFR (%) 180", "Insurance", False),
    COL_SURCHARGE: ("Surcharge", "Others", True),
    COL_TOTAL: ("Total precio ABI", "Total Landing Cost ", True),
}

# Index series for forecast
FORECAST_INDEX_SERIES = "ICIS China Mid"
INDEX_LAG_MONTHS = 4  # N-4

# Front-end canonical columns
FRONT_END_HEADERS = [
    "Data Type",
    "Source File ",
    "Supplier Name",
    "Destination Country",
    "Time_Period ",
    "Time Period Year",
    "Time Period Month",
    "Raw Cost Breakdown",
    "Resin Index Type",
    "Forecast Resin Index Type",
    "Mapping Columns",
    "Column Required for Calculation",
    "Value ",
    "TLC Formula",
]

MONTH_LABELS = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}


# ---------------------------------------------------------------------------
# Date parsing
# ---------------------------------------------------------------------------
def parse_date(val) -> date | None:
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        val = val.strip()
        for fmt in ("%d-%b-%Y", "%d %b %Y", "%Y-%m-%d", "%d-%B-%Y"):
            try:
                return datetime.strptime(val, fmt).date()
            except ValueError:
                continue
    return None


def time_period_label(year: int, month: int) -> str:
    return f"{MONTH_LABELS[month]} {year}"


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------
def extract_monthly_data(source_path: Path) -> list[dict]:
    """Read ICIS Price History sheet, extract rows with data, group by month (last row per month)."""
    wb = load_workbook(source_path, read_only=True, data_only=True)
    ws = wb[SOURCE_SHEET]

    # Collect all rows with index data
    raw_rows = []
    for row_idx in range(DATA_START_ROW, 500):
        index_val = ws.cell(row_idx, COL_INDEX).value
        if index_val is None:
            continue

        date_val = parse_date(ws.cell(row_idx, COL_DATE).value)
        if date_val is None:
            continue

        row_data = {"row": row_idx, "date": date_val}
        for col_idx in COMPONENT_COLUMNS:
            cell_val = ws.cell(row_idx, col_idx).value
            if cell_val is not None and isinstance(cell_val, (int, float)):
                row_data[col_idx] = cell_val
            else:
                row_data[col_idx] = None
        raw_rows.append(row_data)

    wb.close()

    # Group by year-month, take the LAST row per month (latest weekly assessment)
    months: dict[tuple[int, int], list[dict]] = defaultdict(list)
    for row_data in raw_rows:
        key = (row_data["date"].year, row_data["date"].month)
        months[key].append(row_data)

    monthly_data = []
    for (year, month), rows in sorted(months.items()):
        last_row = rows[-1]  # Take last weekly value of the month
        monthly_data.append({
            "year": year,
            "month": month,
            "time_period": time_period_label(year, month),
            "source_row": last_row["row"],
            "date": last_row["date"],
            **{col_idx: last_row.get(col_idx) for col_idx in COMPONENT_COLUMNS},
        })

    return monthly_data


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
def validate_tlc(monthly_data: list[dict]) -> list[dict]:
    """Verify that sum of components matches Total for each month."""
    validation = []
    calc_cols = [COL_INDEX, COL_SC, COL_ADDER, COL_LOGISTICS, COL_LANDED, COL_FINANCE, COL_SURCHARGE]

    for entry in monthly_data:
        total_source = entry.get(COL_TOTAL)
        if total_source is None:
            continue

        computed = sum(entry.get(c) or 0.0 for c in calc_cols)
        diff = abs(total_source - computed)
        status = "match" if diff <= 1.0 else "check"

        validation.append({
            "time_period": entry["time_period"],
            "computed_tlc": round(computed, 4),
            "source_tlc": round(total_source, 4),
            "difference": round(diff, 4),
            "status": status,
        })

    return validation


# ---------------------------------------------------------------------------
# Front-end row building (Actual)
# ---------------------------------------------------------------------------
def build_actual_front_end_rows(monthly_data: list[dict]) -> list[dict]:
    """Build one row per component per month for actual data."""
    rows = []
    source_file_name = SOURCE_FILE.name

    for entry in monthly_data:
        for col_idx, (raw_label, mapping_col, required) in COMPONENT_COLUMNS.items():
            value = entry.get(col_idx)
            if value is None:
                continue

            rows.append({
                "Data Type": "Actual",
                "Source File ": source_file_name,
                "Supplier Name": DEFAULT_METADATA["supplier"],
                "Destination Country": DEFAULT_METADATA["destination_country"],
                "Time_Period ": entry["time_period"],
                "Time Period Year": entry["year"],
                "Time Period Month": entry["month"],
                "Raw Cost Breakdown": raw_label,
                "Resin Index Type": DEFAULT_METADATA["resin_index_type"],
                "Forecast Resin Index Type": None,
                "Mapping Columns": mapping_col,
                "Column Required for Calculation": "Yes" if required else "No",
                "Value ": round(value, 6) if isinstance(value, float) else value,
                "TLC Formula": DEFAULT_METADATA["tlc_formula"],
            })

    return rows


# ---------------------------------------------------------------------------
# Forecast
# ---------------------------------------------------------------------------
def read_index_reference(csv_path: Path, series_name: str) -> dict[tuple[int, int], float]:
    """Read the ICIS resin index reference table, filter to the given series."""
    index_by_period: dict[tuple[int, int], float] = {}

    with open(csv_path, encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("resin_index_type", "").strip() != series_name:
                continue
            try:
                year = int(row["time_period_year"])
                month = int(row["time_period_month"])
                value = float(row["value"])
                index_by_period[(year, month)] = value
            except (ValueError, KeyError):
                continue

    return index_by_period


def apply_lag(year: int, month: int, lag_months: int) -> tuple[int, int]:
    """Shift a period backward by lag_months."""
    total_months = year * 12 + (month - 1) - lag_months
    return total_months // 12, (total_months % 12) + 1


def build_forecast_rows(
    monthly_data: list[dict],
    index_reference: dict[tuple[int, int], float],
    forecast_end: date,
) -> list[dict]:
    """Build forecast rows from the month after last actual through forecast_end."""
    if not monthly_data:
        return []

    # Get latest actual period
    last_entry = monthly_data[-1]
    last_year, last_month = last_entry["year"], last_entry["month"]

    # Get fixed cost components from latest actual period
    latest_adder = last_entry.get(COL_ADDER) or 0.0
    latest_logistics = last_entry.get(COL_LOGISTICS) or 0.0
    latest_landed = last_entry.get(COL_LANDED) or 0.0
    latest_sc = last_entry.get(COL_SC) or 0.0
    latest_surcharge = last_entry.get(COL_SURCHARGE) or 0.0

    # Calculate finance rate from latest period
    latest_index = last_entry.get(COL_INDEX) or 0.0
    base_for_finance = latest_index + latest_adder + latest_logistics + latest_landed
    if base_for_finance > 0 and last_entry.get(COL_FINANCE):
        finance_rate = last_entry[COL_FINANCE] / base_for_finance
    else:
        finance_rate = 0.0375  # default fallback ~3.75%

    source_file_name = SOURCE_FILE.name
    forecast_rows = []

    # Generate forecast months
    current_year, current_month = last_year, last_month
    while True:
        # Advance to next month
        current_month += 1
        if current_month > 12:
            current_month = 1
            current_year += 1

        if date(current_year, current_month, 1) > forecast_end:
            break

        # Get lagged index value
        lagged_year, lagged_month = apply_lag(current_year, current_month, INDEX_LAG_MONTHS)
        index_value = index_reference.get((lagged_year, lagged_month))

        if index_value is None:
            # Try without lag if lagged period not available
            index_value = index_reference.get((current_year, current_month))
            if index_value is None:
                continue

        # Calculate finance fee
        base = index_value + latest_adder + latest_logistics + latest_landed
        finance_fee = base * finance_rate

        # Calculate total
        total = index_value + latest_sc + latest_adder + latest_logistics + latest_landed + finance_fee + latest_surcharge

        time_period = time_period_label(current_year, current_month)

        # Build component rows for this forecast month
        forecast_components = [
            (COL_INDEX, index_value),
            (COL_SC, latest_sc),
            (COL_ADDER, latest_adder),
            (COL_LOGISTICS, latest_logistics),
            (COL_LANDED, latest_landed),
            (COL_FINANCE, finance_fee),
            (COL_SURCHARGE, latest_surcharge),
            (COL_TOTAL, total),
        ]

        for col_idx, value in forecast_components:
            if value is None or value == 0.0:
                # Skip zero components that were never present
                if col_idx in (COL_SC, COL_SURCHARGE) and last_entry.get(col_idx) is None:
                    continue

            raw_label, mapping_col, required = COMPONENT_COLUMNS[col_idx]
            forecast_rows.append({
                "Data Type": "Forecast",
                "Source File ": source_file_name,
                "Supplier Name": DEFAULT_METADATA["supplier"],
                "Destination Country": DEFAULT_METADATA["destination_country"],
                "Time_Period ": time_period,
                "Time Period Year": current_year,
                "Time Period Month": current_month,
                "Raw Cost Breakdown": raw_label,
                "Resin Index Type": DEFAULT_METADATA["resin_index_type"],
                "Forecast Resin Index Type": f"{FORECAST_INDEX_SERIES} (N-{INDEX_LAG_MONTHS})",
                "Mapping Columns": mapping_col,
                "Column Required for Calculation": "Yes" if required else "No",
                "Value ": round(value, 6) if isinstance(value, float) else value,
                "TLC Formula": DEFAULT_METADATA["tlc_formula"],
            })

    return forecast_rows


# ---------------------------------------------------------------------------
# Output writing
# ---------------------------------------------------------------------------
def write_front_end_excel(path: Path, rows: list[dict]) -> None:
    """Write the front-end standardized Excel file."""
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("front_end_standardized")

    # Write headers
    ws.append(FRONT_END_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    # Write data
    for row in rows:
        ws.append([row.get(h) for h in FRONT_END_HEADERS])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Auto-size columns
    for col_idx, header in enumerate(FRONT_END_HEADERS, start=1):
        width = min(max(len(header) + 2, 12), 45)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(path)


def write_front_end_csv(path: Path, rows: list[dict]) -> None:
    """Write the front-end standardized CSV file."""
    path.parent.mkdir(parents=True, exist_ok=True)

    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=FRONT_END_HEADERS)
        writer.writeheader()
        writer.writerows(rows)


def write_validation_json(path: Path, validation: list[dict], metadata: dict) -> None:
    """Write validation summary."""
    path.parent.mkdir(parents=True, exist_ok=True)
    output = {
        "pipeline": PIPELINE_NAME,
        "run_timestamp": datetime.now().isoformat(timespec="seconds"),
        "metadata": metadata,
        "validation_rows": len(validation),
        "all_match": all(v["status"] == "match" for v in validation),
        "details": validation,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, default=str)


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description="Run El Salvador & Honduras Sell Side pipeline")
    parser.add_argument("--skip-forecast", action="store_true", help="Skip forecast generation")
    parser.add_argument("--forecast-end", type=str, default="2026-12-01",
                        help="Forecast end date (YYYY-MM-DD)")
    args = parser.parse_args()

    forecast_end = datetime.strptime(args.forecast_end, "%Y-%m-%d").date()

    print(f"{'='*60}")
    print(f"Pipeline: {PIPELINE_NAME}")
    print(f"Source: {SOURCE_FILE.name}")
    print(f"Destination: {DEFAULT_METADATA['destination_country']}")
    print(f"{'='*60}")

    # --- Step 1: Extract monthly data ---
    print("\n[1/5] Extracting monthly data from ICIS Price History...")
    if not SOURCE_FILE.exists():
        print(f"ERROR: Source file not found: {SOURCE_FILE}")
        sys.exit(1)

    monthly_data = extract_monthly_data(SOURCE_FILE)
    print(f"  Extracted {len(monthly_data)} monthly periods")
    print(f"  Period range: {monthly_data[0]['time_period']} to {monthly_data[-1]['time_period']}")

    # --- Step 2: Validate TLC ---
    print("\n[2/5] Validating TLC calculations...")
    validation = validate_tlc(monthly_data)
    match_count = sum(1 for v in validation if v["status"] == "match")
    print(f"  Validated {len(validation)} periods: {match_count} match, {len(validation) - match_count} check")

    # --- Step 3: Build actual front-end rows ---
    print("\n[3/5] Building actual front-end rows...")
    actual_rows = build_actual_front_end_rows(monthly_data)
    print(f"  Generated {len(actual_rows)} actual rows")

    # --- Step 4: Build forecast rows ---
    forecast_rows = []
    if not args.skip_forecast:
        print("\n[4/5] Building forecast rows...")
        if INDEX_REFERENCE_CSV.exists():
            index_reference = read_index_reference(INDEX_REFERENCE_CSV, FORECAST_INDEX_SERIES)
            print(f"  Loaded {len(index_reference)} index reference values for '{FORECAST_INDEX_SERIES}'")

            forecast_rows = build_forecast_rows(monthly_data, index_reference, forecast_end)
            print(f"  Generated {len(forecast_rows)} forecast rows")
            if forecast_rows:
                first_fc = forecast_rows[0]["Time_Period "]
                last_fc = forecast_rows[-1]["Time_Period "]
                print(f"  Forecast range: {first_fc} to {last_fc}")
        else:
            print(f"  WARNING: Index reference not found at {INDEX_REFERENCE_CSV}")
            print("  Skipping forecast.")
    else:
        print("\n[4/5] Skipping forecast (--skip-forecast)")

    # --- Step 5: Write outputs ---
    print("\n[5/5] Writing outputs...")

    all_rows = actual_rows + forecast_rows

    # Front-end standardized Excel
    excel_path = FRONT_END_DIR / f"{PIPELINE_NAME}_actual_forecast_front_end_standardized.xlsx"
    write_front_end_excel(excel_path, all_rows)
    print(f"  Excel: {excel_path.relative_to(REPO_ROOT)}")

    # Front-end standardized CSV
    csv_path = FRONT_END_DIR / f"{PIPELINE_NAME}_actual_forecast_front_end_standardized.csv"
    write_front_end_csv(csv_path, all_rows)
    print(f"  CSV: {csv_path.relative_to(REPO_ROOT)}")

    # Validation JSON
    validation_path = VALIDATION_DIR / "validation_summary.json"
    write_validation_json(validation_path, validation, {
        **DEFAULT_METADATA,
        "source_file": SOURCE_FILE.name,
        "source_sheet": SOURCE_SHEET,
        "actual_periods": len(monthly_data),
        "actual_rows": len(actual_rows),
        "forecast_rows": len(forecast_rows),
        "total_rows": len(all_rows),
    })
    print(f"  Validation: {validation_path.relative_to(REPO_ROOT)}")

    # --- Summary ---
    print(f"\n{'='*60}")
    print(f"PIPELINE COMPLETE")
    print(f"  Total front-end rows: {len(all_rows)}")
    print(f"  Actual: {len(actual_rows)} | Forecast: {len(forecast_rows)}")
    print(f"  Output: {excel_path.relative_to(REPO_ROOT)}")
    print(f"{'='*60}")


if __name__ == "__main__":
    main()
