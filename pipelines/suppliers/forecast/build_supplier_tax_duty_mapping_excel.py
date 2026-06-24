from __future__ import annotations

from pathlib import Path
import re

import openpyxl
import pandas as pd

BASE = Path(__file__).resolve().parents[3]
SRC_CSV = (
    BASE
    / "data"
    / "processed"
    / "current"
    / "supplier"
    / "Data_Standardized_Front_End_Data_Model.csv"
)
OUT_XLSX = (
    BASE
    / "data"
    / "processed"
    / "current"
    / "supplier"
    / "All_Suppliers_Destinations_Taxes_Duties_Derivation.xlsx"
)

TAX_DUTY_KEYWORDS = re.compile(r"tax|duty|import|aduana|arancel|impuesto", re.IGNORECASE)
MAPPED_TAX_COLUMNS = {
    "tax",
    "customs clearance",
    "internalization",
}
SUM_RANGE_RE = re.compile(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)(.*)", re.IGNORECASE)


def classify_constant_or_variable(df: pd.DataFrame) -> str:
    values = (
        pd.to_numeric(df["Value "], errors="coerce")
        .dropna()
        .astype(float)
        .round(8)
        .unique()
        .tolist()
    )
    return "Constant" if len(values) <= 1 else "Variable"


def is_tax_duty_row(raw_label: str, mapped_label: str) -> bool:
    raw = (raw_label or "").strip()
    mapped = (mapped_label or "").strip().lower()
    if mapped in MAPPED_TAX_COLUMNS:
        return True
    return bool(TAX_DUTY_KEYWORDS.search(raw))


def first_non_empty(series: pd.Series) -> str:
    for value in series.fillna("").astype(str):
        trimmed = value.strip()
        if trimmed:
            return trimmed
    return ""


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def prettify_formula(formula: str) -> str:
    clean = (formula or "").strip()
    if clean.startswith("="):
        clean = clean[1:]
    clean = clean.replace("$", "")

    match = SUM_RANGE_RE.fullmatch(clean)
    if match:
        left, right, suffix = match.groups()
        return f"({left}+{right}){suffix}"

    return clean


def resolve_source_file(source_file: str, file_index: dict[str, Path]) -> Path | None:
    source = (source_file or "").strip()
    if not source:
        return None

    source_path = Path(source)
    if source_path.is_absolute() and source_path.exists():
        return source_path

    direct = BASE / source_path
    if direct.exists():
        return direct

    by_name = file_index.get(source_path.name.lower())
    if by_name and by_name.exists():
        return by_name

    return None


def extract_derived_formula_from_workbook(
    workbook_path: Path,
    raw_label: str,
    formula_cache: dict[tuple[str, str], str],
) -> str:
    cache_key = (str(workbook_path), normalize_text(raw_label))
    if cache_key in formula_cache:
        return formula_cache[cache_key]

    target = normalize_text(raw_label)
    if not target:
        formula_cache[cache_key] = ""
        return ""

    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=False)
    except Exception:
        formula_cache[cache_key] = ""
        return ""

    found_formulas: list[str] = []

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            label_found = False
            for cell in row:
                if isinstance(cell.value, str) and normalize_text(cell.value) == target:
                    label_found = True
                    break

            if not label_found:
                continue

            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    pretty = prettify_formula(cell.value)
                    if pretty:
                        found_formulas.append(pretty)

    unique: list[str] = []
    seen: set[str] = set()
    for formula in found_formulas:
        key = formula.lower()
        if key in seen:
            continue
        seen.add(key)
        unique.append(formula)

    result = unique[0] if unique else ""
    formula_cache[cache_key] = result
    return result


def build_file_index(source_files: pd.Series) -> dict[str, Path]:
    names = {
        Path(value).name.lower()
        for value in source_files.dropna().astype(str)
        if str(value).strip()
    }

    index: dict[str, Path] = {}
    supplier_root = BASE / "data" / "raw" / "suppliers"

    for name in names:
        matches = list(supplier_root.rglob(name))
        if matches:
            index[name] = matches[0]
            continue

        matches = list(BASE.rglob(name))
        if matches:
            index[name] = matches[0]

    return index


def main() -> None:
    df = pd.read_csv(SRC_CSV)

    actual_df = df[df["Data Type"].astype(str).str.lower() == "actual"].copy()
    actual_df = actual_df[
        actual_df.apply(
            lambda row: is_tax_duty_row(
                str(row.get("Raw Cost Breakdown", "")),
                str(row.get("Mapping Columns", "")),
            ),
            axis=1,
        )
    ]

    if actual_df.empty:
        raise ValueError("No tax/duty/import rows found in actual supplier data")

    file_index = build_file_index(df.get("Source File ", pd.Series(dtype=str)))
    formula_cache: dict[tuple[str, str], str] = {}

    rows: list[dict[str, str]] = []
    grouped = actual_df.groupby(
        ["Supplier Name", "Destination Country", "Raw Cost Breakdown", "Mapping Columns"],
        dropna=False,
    )

    for (supplier, destination, raw_label, mapped_label), group in grouped:
        classification = classify_constant_or_variable(group)
        source_file = first_non_empty(group.get("Source File ", pd.Series(dtype=str)))

        derived_formula = ""
        if classification == "Variable":
            resolved_path = resolve_source_file(source_file, file_index)
            if resolved_path is not None:
                derived_formula = extract_derived_formula_from_workbook(
                    resolved_path,
                    str(raw_label),
                    formula_cache,
                )

            if not derived_formula:
                derived_formula = (
                    f"Refer supplier sheet: {source_file}"
                    if source_file
                    else "Refer supplier sheet"
                )

        rows.append(
            {
                "SupplierName": str(supplier),
                "Destination COuntry": str(destination),
                "Raw Taxes & Duties(from supplier sheet)": str(raw_label),
                "Mapped Taxes & Duties (what we have mapped )": str(mapped_label),
                "Constant or Variable": classification,
                "Derived Formula": derived_formula,
            }
        )

    out_df = pd.DataFrame(rows).sort_values(
        by=[
            "SupplierName",
            "Destination COuntry",
            "Mapped Taxes & Duties (what we have mapped )",
            "Raw Taxes & Duties(from supplier sheet)",
        ]
    )

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        out_df.to_excel(writer, sheet_name="tax_duty_mapping", index=False)

    print(f"Created: {OUT_XLSX}")
    print(f"Rows exported: {len(out_df)}")
    print(
        out_df[
            [
                "SupplierName",
                "Destination COuntry",
                "Raw Taxes & Duties(from supplier sheet)",
                "Mapped Taxes & Duties (what we have mapped )",
                "Constant or Variable",
                "Derived Formula",
            ]
        ]
        .head(20)
        .to_string(index=False)
    )


if __name__ == "__main__":
    main()
