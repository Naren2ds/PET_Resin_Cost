"""Compare Supplier resin-index lags with all MR source-country lags.

For the same destination/year/month/data type, every Supplier lag is compared
with every available Market Research source country. Explicit M-n and N-n
notations are compared numerically; a label without lag notation is lag 0.

Run from PET_Resin_Cost:
    python testing_scripts/validate_supplier_market_lags.py --destination Colombia --year 2026
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUPPLIER_INPUT = PROJECT_ROOT / "data/processed/current/supplier/Data_Standardized_Front_End_Data_Model.xlsx"
MR_INPUT = PROJECT_ROOT / "data/processed/current/market_research/Data_Standardized_MR_Front_End_Data_Model.xlsx"
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")


def clean(value: Any) -> str:
    text = "" if value is None else str(value).replace("\xa0", " ").strip()
    return "" if text.casefold() in {"nan", "none", "nat"} else text


def read_rows(path: Path) -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["front_end_data_model"] if "front_end_data_model" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(values)]
        for row in values:
            if any(value is not None and clean(value) for value in row):
                yield dict(zip(headers, row))
    finally:
        workbook.close()


def is_resin_index(row: dict[str, Any]) -> bool:
    return clean(row.get("Mapping Columns")).casefold() in {"resin index vpet", "index", "seguro"}


def period(row: dict[str, Any]) -> tuple[int, int] | None:
    try:
        year = int(float(clean(row.get("Time Period Year"))))
        month = int(float(clean(row.get("Time Period Month"))))
        return (year, month) if 1 <= month <= 12 else None
    except ValueError:
        return None


def active_index_label(row: dict[str, Any]) -> str:
    data_type = clean(row.get("Data Type")).casefold()
    actual = clean(row.get("Resin Index Type"))
    forecast = clean(row.get("Forecast Resin Index Type"))
    return (forecast or actual) if data_type == "forecast" else actual


def extract_lag(label: str) -> tuple[int, str]:
    """Return numeric lag and its notation; unmarked labels mean current month."""
    match = re.search(r"\b([MN])\s*-\s*(\d+)\b", label, flags=re.IGNORECASE)
    if match:
        return int(match.group(2)), f"{match.group(1).upper()}-{int(match.group(2))}"
    return 0, "Current month (0)"


def collect(rows: Iterable[dict[str, Any]], destination: str | None, year: int | None):
    result = defaultdict(lambda: defaultdict(set))
    for row in rows:
        if not is_resin_index(row):
            continue
        entity = clean(row.get("Supplier Name"))
        market = clean(row.get("Destination Country"))
        data_type = clean(row.get("Data Type")).title()
        row_period = period(row)
        label = active_index_label(row)
        if not entity or not market or data_type not in {"Actual", "Forecast"} or not row_period or not label:
            continue
        row_year, month = row_period
        if destination and market.casefold() != destination.casefold():
            continue
        if year is not None and row_year != year:
            continue
        lag, notation = extract_lag(label)
        result[(market, row_year, month, data_type, entity)][lag].add((label, notation))
    return result


def build_comparisons(supplier_data, mr_data) -> list[dict[str, Any]]:
    mr_by_period = defaultdict(list)
    for (market, year, month, data_type, source), lag_groups in mr_data.items():
        mr_by_period[(market.casefold(), year, month, data_type)].append((source, lag_groups))
    comparisons = []
    for (market, year, month, data_type, supplier), supplier_groups in sorted(supplier_data.items()):
        mr_groups = mr_by_period.get((market.casefold(), year, month, data_type), [])
        if not mr_groups:
            comparisons.append({"market": market, "year": year, "month": month, "data_type": data_type,
                                "supplier": supplier, "supplier_groups": supplier_groups, "mr_source": "No MR data",
                                "mr_groups": {}, "status": "FAIL", "note": "No matching MR source-country data"})
            continue
        for source, market_groups in sorted(mr_groups):
            supplier_lags = set(supplier_groups)
            market_lags = set(market_groups)
            comparisons.append({"market": market, "year": year, "month": month, "data_type": data_type,
                                "supplier": supplier, "supplier_groups": supplier_groups, "mr_source": source,
                                "mr_groups": market_groups,
                                "status": "PASS" if supplier_lags == market_lags and len(supplier_lags) == 1 else "FAIL",
                                "note": "" if supplier_lags == market_lags and len(supplier_lags) == 1 else "Supplier and MR numeric lags differ"})
    return comparisons


def group_text(groups) -> tuple[str, str, str]:
    if not groups:
        return "", "", ""
    lags = ", ".join(str(value) for value in sorted(groups))
    notations = sorted({notation for items in groups.values() for _, notation in items})
    labels = sorted({label for items in groups.values() for label, _ in items})
    return lags, "; ".join(notations), "; ".join(labels)


def write_excel(output: Path, rows: list[dict[str, Any]], supplier_input: Path, mr_input: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    details = workbook.create_sheet("Lag Comparisons")
    blue = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    summary.append(["Metric", "Value"])
    for item in (("Supplier Input", str(supplier_input)), ("MR Input", str(mr_input)),
                 ("Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                 ("Comparisons", len(rows)), ("Passed", sum(r["status"] == "PASS" for r in rows)),
                 ("Failed", sum(r["status"] == "FAIL" for r in rows)),
                 ("Overall Status", "PASS" if rows and all(r["status"] == "PASS" for r in rows) else "FAIL"),
                 ("Lag Rule", "M-n and N-n compare by numeric n; no notation = current-month lag 0")):
        summary.append(item)
    headers = ["Destination", "Year", "Month", "Data Type", "Supplier", "Supplier Lag",
               "Supplier Notation", "Supplier Index Label", "MR Source Country", "MR Lag",
               "MR Notation", "MR Index Label", "Status", "Notes"]
    details.append(headers)
    for row in rows:
        supplier_lag, supplier_notation, supplier_label = group_text(row["supplier_groups"])
        mr_lag, mr_notation, mr_label = group_text(row["mr_groups"])
        details.append([row["market"], row["year"], MONTHS[row["month"] - 1], row["data_type"],
                        row["supplier"], supplier_lag, supplier_notation, supplier_label, row["mr_source"],
                        mr_lag, mr_notation, mr_label, row["status"], row["note"]])
        status_cell = details.cell(details.max_row, 13)
        status_cell.fill = green if row["status"] == "PASS" else red
        status_cell.font = Font(bold=True)
    for sheet in (summary, details):
        for cell in sheet[1]:
            cell.fill = blue; cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    summary.column_dimensions["A"].width = 24; summary.column_dimensions["B"].width = 110
    for index, width in enumerate([24, 9, 12, 12, 34, 14, 20, 55, 24, 12, 20, 55, 12, 38], 1):
        details.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare Supplier resin-index lags with every MR source country.")
    parser.add_argument("--destination", help="Validate one destination (recommended)")
    parser.add_argument("--year", type=int, help="Validate one year")
    parser.add_argument("--supplier-input", type=Path, default=SUPPLIER_INPUT)
    parser.add_argument("--mr-input", type=Path, default=MR_INPUT)
    parser.add_argument("--output-excel", type=Path)
    args = parser.parse_args()
    output = args.output_excel or Path(f"supplier_market_lag_validation_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    for path in (args.supplier_input, args.mr_input):
        if not path.is_file():
            print(f"ERROR: Input not found: {path}", file=sys.stderr); return 2
    try:
        supplier = collect(read_rows(args.supplier_input), args.destination, args.year)
        market = collect(read_rows(args.mr_input), args.destination, args.year)
        rows = build_comparisons(supplier, market)
        if not rows:
            print("ERROR: No Supplier/MR comparisons matched the filters.", file=sys.stderr); return 2
        write_excel(output, rows, args.supplier_input, args.mr_input)
    except (OSError, ValueError, ImportError, StopIteration) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); return 2
    passed = sum(row["status"] == "PASS" for row in rows)
    print(f"Comparisons: {len(rows)} | Passed: {passed} | Failed: {len(rows)-passed}")
    print(f"Excel report: {output.resolve()}")
    return 0 if passed == len(rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
