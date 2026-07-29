"""Validate TrendsPage Actual/Forecast data coverage and rendered line styles.

Creates month-wise Excel evidence for every Supplier and Market Research source
country, separately for TLC and Resin Index.

Run from PET_Resin_Cost:
    python testing_scripts/test_trends_page_line_styles.py --year 2026
"""
from __future__ import annotations

import argparse
import re
import sys
import unittest
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TRENDS_PAGE = PROJECT_ROOT / "apps/frontend/src/pages/TrendsPage.tsx"
SUPPLIER_INPUT = PROJECT_ROOT / "data/processed/current/supplier/Data_Standardized_Front_End_Data_Model.xlsx"
MR_INPUT = PROJECT_ROOT / "data/processed/current/market_research/Data_Standardized_MR_Front_End_Data_Model.xlsx"
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")

STYLE_KEYS = {
    ("Supplier", "TLC", "Actual"): "supplierActual",
    ("Supplier", "TLC", "Forecast"): "supplierForecast",
    ("Market Research", "TLC", "Actual"): "market_${index}_actual",
    ("Market Research", "TLC", "Forecast"): "market_${index}_forecast",
    ("Supplier", "Resin Index", "Actual"): "supplierIndexActual",
    ("Supplier", "Resin Index", "Forecast"): "supplierIndexForecast",
    ("Market Research", "Resin Index", "Actual"): "marketIndex_${index}_actual",
    ("Market Research", "Resin Index", "Forecast"): "marketIndex_${index}_forecast",
}
EXPECTED_ACTUAL_KEYS = {key for (dataset, metric, kind), key in STYLE_KEYS.items() if kind == "Actual"}
EXPECTED_FORECAST_KEYS = {key for (dataset, metric, kind), key in STYLE_KEYS.items() if kind == "Forecast"}


def clean(value: Any) -> str:
    text = "" if value is None else str(value).replace("\xa0", " ").strip()
    return "" if text.casefold() in {"nan", "none", "nat"} else text


def line_elements(source: str) -> dict[str, str]:
    elements = {}
    for match in re.finditer(r"<Line\b(?P<attributes>.*?)/>", source, re.DOTALL):
        element = match.group(0)
        key_match = re.search(r'\bdataKey\s*=\s*(?:"(?P<quoted>[^"]+)"|\{`(?P<template>[^`]+)`\})', element)
        if key_match:
            key = key_match.group("quoted") or key_match.group("template")
            elements[key] = element
    return elements


def dash_pattern(element: str) -> str | None:
    match = re.search(r'\bstrokeDasharray\s*=\s*["\']([^"\']+)["\']', element)
    return match.group(1).strip() if match else None


def rendered_style(lines: dict[str, str], key: str) -> tuple[str, str, bool]:
    element = lines.get(key)
    if element is None:
        return "Missing series", "", False
    pattern = dash_pattern(element)
    if pattern:
        valid = bool(re.fullmatch(r"\d+(?:[ ,]+\d+)+", pattern))
        return "Dotted", pattern, valid
    return "Solid", "", True


def read_rows(path: Path) -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["front_end_data_model"] if "front_end_data_model" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(values)]
        for row in values:
            if any(value is not None and clean(value) for value in row):
                yield dict(zip(headers, row))
    finally:
        workbook.close()


def parse_period(row: dict[str, Any]) -> tuple[int, int] | None:
    try:
        year = int(float(clean(row.get("Time Period Year"))))
        month = int(float(clean(row.get("Time Period Month"))))
        if 1 <= month <= 12:
            return year, month
    except ValueError:
        pass
    value = row.get("Time_Period")
    if isinstance(value, (date, datetime)):
        return value.year, value.month
    for fmt in ("%Y-%m-%d", "%Y-%m", "%B %Y", "%b %Y"):
        try:
            parsed = datetime.strptime(clean(value), fmt)
            return parsed.year, parsed.month
        except ValueError:
            pass
    return None


def metric_for_row(row: dict[str, Any]) -> str | None:
    mapping = clean(row.get("Mapping Columns")).casefold()
    if mapping.startswith("total landing cost"):
        return "TLC"
    if mapping in {"resin index vpet", "index", "seguro"}:
        return "Resin Index"
    return None


def build_coverage(rows: Iterable[dict[str, Any]], year_filter: int | None,
                   entity_filter: str | None, market_filter: str | None):
    coverage = defaultdict(lambda: defaultdict(set))
    for row in rows:
        metric = metric_for_row(row)
        entity = clean(row.get("Supplier Name"))
        market = clean(row.get("Destination Country"))
        data_type = clean(row.get("Data Type")).title()
        period = parse_period(row)
        if not metric or not entity or not market or data_type not in {"Actual", "Forecast"} or not period:
            continue
        year, month = period
        if year_filter is not None and year != year_filter:
            continue
        if entity_filter and entity.casefold() != entity_filter.casefold():
            continue
        if market_filter and market.casefold() != market_filter.casefold():
            continue
        coverage[(entity, market, year, metric)][month].add(data_type)
    return coverage


def report_rows(dataset: str, coverage, lines: dict[str, str]) -> list[dict[str, Any]]:
    rows = []
    for (entity, market, year, metric), month_types in sorted(coverage.items(), key=lambda item: (item[0][2], item[0][1], item[0][0], item[0][3])):
        actual_style, actual_dash, actual_valid = rendered_style(lines, STYLE_KEYS[(dataset, metric, "Actual")])
        forecast_style, forecast_dash, forecast_valid = rendered_style(lines, STYLE_KEYS[(dataset, metric, "Forecast")])
        month_values = []
        actual_months, forecast_months = set(), set()
        for month in range(1, 13):
            types = month_types.get(month, set())
            values = []
            if "Actual" in types:
                actual_months.add(month)
                values.append(f"Actual | {actual_style}")
            if "Forecast" in types:
                forecast_months.add(month)
                values.append(f"Forecast | {forecast_style}")
            month_values.append("; ".join(values) if values else "Missing")
        data_ok = len(actual_months | forecast_months) == 12
        style_ok = actual_style == "Solid" and forecast_style == "Dotted" and actual_valid and forecast_valid
        rows.append({
            "dataset": dataset, "entity": entity, "market": market, "year": year, "metric": metric,
            "months": month_values, "actual_count": len(actual_months), "forecast_count": len(forecast_months),
            "actual_months": ", ".join(MONTHS[m - 1] for m in sorted(actual_months)),
            "forecast_months": ", ".join(MONTHS[m - 1] for m in sorted(forecast_months)),
            "actual_style": actual_style, "forecast_style": forecast_style,
            "actual_dash": actual_dash or "None", "forecast_dash": forecast_dash or "None",
            "coverage_status": "PASS" if data_ok else "FAIL", "style_status": "PASS" if style_ok else "FAIL",
            "status": "PASS" if data_ok and style_ok else "FAIL",
        })
    return rows


def write_excel(output: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    blue = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    actual_fill = PatternFill("solid", fgColor="D9EAF7")
    forecast_fill = PatternFill("solid", fgColor="FFF2CC")
    missing_fill = PatternFill("solid", fgColor="F4CCCC")
    summary.append(["Dataset", "Metric", "Rows", "Passed", "Failed"])
    for dataset in ("Supplier", "Market Research"):
        for metric in ("TLC", "Resin Index"):
            selected = [row for row in rows if row["dataset"] == dataset and row["metric"] == metric]
            summary.append([dataset, metric, len(selected), sum(r["status"] == "PASS" for r in selected), sum(r["status"] == "FAIL" for r in selected)])
    summary.append(["TOTAL", "", len(rows), sum(r["status"] == "PASS" for r in rows), sum(r["status"] == "FAIL" for r in rows)])

    for dataset, sheet_name, entity_header in (("Supplier", "Supplier Trends", "Supplier"), ("Market Research", "MR Trends", "MR Source Country")):
        sheet = workbook.create_sheet(sheet_name)
        headers = [entity_header, "Destination Market", "Year", "Metric", *MONTHS,
                   "Actual Month Count", "Forecast Month Count", "Actual Months", "Forecast Months",
                   "Actual Style", "Actual Dash", "Forecast Style", "Forecast Dash",
                   "Coverage Status", "Style Status", "Overall Status"]
        sheet.append(headers)
        for row in (item for item in rows if item["dataset"] == dataset):
            sheet.append([row["entity"], row["market"], row["year"], row["metric"], *row["months"],
                          row["actual_count"], row["forecast_count"], row["actual_months"], row["forecast_months"],
                          row["actual_style"], row["actual_dash"], row["forecast_style"], row["forecast_dash"],
                          row["coverage_status"], row["style_status"], row["status"]])
            current = sheet.max_row
            for column in range(5, 17):
                cell = sheet.cell(current, column)
                if cell.value == "Missing": cell.fill = missing_fill
                elif str(cell.value).startswith("Actual"): cell.fill = actual_fill
                elif str(cell.value).startswith("Forecast"): cell.fill = forecast_fill
            for column in range(len(headers) - 2, len(headers) + 1):
                cell = sheet.cell(current, column)
                cell.fill = green if cell.value == "PASS" else red
                cell.font = Font(bold=True)
        sheet.freeze_panes = "E2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [32, 24, 9, 14, *([22] * 12), 18, 20, 55, 55, 14, 14, 16, 16, 16, 14, 16]
        for index, width in enumerate(widths, 1): sheet.column_dimensions[get_column_letter(index)].width = width
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = blue; cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = sheet.freeze_panes or "A2"
        sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output)


class TrendsPageLineStyleTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lines = line_elements(TRENDS_PAGE.read_text(encoding="utf-8"))

    def test_actual_series_are_solid(self):
        for key in EXPECTED_ACTUAL_KEYS:
            with self.subTest(dataKey=key):
                self.assertIn(key, self.lines)
                self.assertIsNone(dash_pattern(self.lines[key]))

    def test_forecast_series_are_dotted(self):
        for key in EXPECTED_FORECAST_KEYS:
            with self.subTest(dataKey=key):
                self.assertIn(key, self.lines)
                self.assertRegex(dash_pattern(self.lines[key]) or "", r"^\d+(?:[ ,]+\d+)+$")


def main() -> int:
    parser = argparse.ArgumentParser(description="Export supplier/MR month-wise Actual/Forecast TrendsPage validation.")
    parser.add_argument("--supplier-input", type=Path, default=SUPPLIER_INPUT)
    parser.add_argument("--mr-input", type=Path, default=MR_INPUT)
    parser.add_argument("--year", type=int, help="Validate only this year")
    parser.add_argument("--supplier", help="Filter Supplier Name/source country")
    parser.add_argument("--market", help="Filter Destination Country")
    parser.add_argument("--output-excel", type=Path)
    args = parser.parse_args()
    output = args.output_excel or Path(f"trends_month_style_validation_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    for path in (args.supplier_input, args.mr_input, TRENDS_PAGE):
        if not path.is_file():
            print(f"ERROR: File not found: {path}", file=sys.stderr); return 2
    test_result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(TrendsPageLineStyleTest))
    try:
        lines = line_elements(TRENDS_PAGE.read_text(encoding="utf-8"))
        supplier = build_coverage(read_rows(args.supplier_input), args.year, args.supplier, args.market)
        mr = build_coverage(read_rows(args.mr_input), args.year, args.supplier, args.market)
        rows = report_rows("Supplier", supplier, lines) + report_rows("Market Research", mr, lines)
        if not rows:
            print("ERROR: No data matched the filters.", file=sys.stderr); return 2
        write_excel(output, rows)
    except (OSError, ValueError, ImportError, StopIteration) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); return 2
    print(f"Rows checked: {len(rows)} | Passed: {sum(r['status']=='PASS' for r in rows)} | Failed: {sum(r['status']=='FAIL' for r in rows)}")
    print(f"Excel report: {output.resolve()}")
    return 0 if test_result.wasSuccessful() and all(row["status"] == "PASS" for row in rows) else 1


if __name__ == "__main__":
    raise SystemExit(main())
