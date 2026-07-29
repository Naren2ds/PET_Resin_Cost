"""Validate 12-month coverage in Supplier and Market Research Excel models.

Run from PET_Resin_Cost:
    python testing_scripts/validate_supplier_market_months.py --year 2026
    python testing_scripts/validate_supplier_market_months.py --output-excel report.xlsx
"""
from __future__ import annotations

import argparse
import csv
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUPPLIER_INPUT = PROJECT_ROOT / "data/processed/current/supplier/Data_Standardized_Front_End_Data_Model.xlsx"
MR_INPUT = PROJECT_ROOT / "data/processed/current/market_research/Data_Standardized_MR_Front_End_Data_Model.xlsx"
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
EXPECTED = set(range(1, 13))
Coverage = dict[tuple[str, str, int], set[int]]


def clean(value: Any) -> str:
    text = "" if value is None else str(value).replace("\xa0", " ").strip()
    return "" if text.casefold() in {"nan", "none", "nat"} else text


def read_rows(path: Path) -> Iterable[dict[str, Any]]:
    if path.suffix.casefold() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            yield from csv.DictReader(handle)
        return
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Inputs must be CSV, XLSX, or XLSM files")
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["front_end_data_model"] if "front_end_data_model" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows)]
        for values in rows:
            if any(value is not None and clean(value) for value in values):
                yield dict(zip(headers, values))
    finally:
        workbook.close()


def parse_period(row: dict[str, Any]) -> tuple[int, int] | None:
    try:
        year = int(float(clean(row.get("Time Period Year"))))
        month = int(float(clean(row.get("Time Period Month"))))
        if 1 <= month <= 12:
            return year, month
    except ValueError:
        pass
    value = row.get("Time_Period")
    if isinstance(value, (date, datetime)):
        return value.year, value.month
    for fmt in ("%Y-%m-%d", "%Y-%m", "%B %Y", "%b %Y"):
        try:
            parsed = datetime.strptime(clean(value), fmt)
            return parsed.year, parsed.month
        except ValueError:
            pass
    return None


def validate(rows: Iterable[dict[str, Any]], year_filter: int | None,
             entity_filter: str | None, market_filter: str | None) -> tuple[Coverage, int]:
    coverage: Coverage = defaultdict(set)
    invalid = 0
    for row in rows:
        entity = clean(row.get("Supplier Name"))
        market = clean(row.get("Destination Country"))
        if not entity or not market:
            continue
        if entity_filter and entity.casefold() != entity_filter.casefold():
            continue
        if market_filter and market.casefold() != market_filter.casefold():
            continue
        period = parse_period(row)
        if period is None:
            invalid += 1
            continue
        year, month = period
        if year_filter is None or year == year_filter:
            coverage[(entity, market, year)].add(month)
    return coverage, invalid


def passed(present: set[int]) -> bool:
    return present == EXPECTED


def write_excel(output: Path, supplier: Coverage, mr: Coverage,
                supplier_input: Path, mr_input: Path,
                supplier_invalid: int, mr_invalid: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    blue = PatternFill("solid", fgColor="1F4E78")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")

    def coverage_sheet(name: str, entity_header: str, data: Coverage) -> None:
        sheet = workbook.create_sheet(name)
        headers = [entity_header, "Destination Market", "Year", *MONTHS, "Months Present", "Missing Months", "Status"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill, cell.font = blue, Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for (entity, market, year), present in sorted(data.items(), key=lambda item: (item[0][2], item[0][1], item[0][0])):
            missing = sorted(EXPECTED - present)
            result = "PASS" if passed(present) else "FAIL"
            sheet.append([entity, market, year, *("Yes" if m in present else "No" for m in range(1, 13)),
                          len(present), ", ".join(MONTHS[m - 1] for m in missing), result])
            row = sheet.max_row
            for column in range(4, 16):
                cell = sheet.cell(row, column)
                cell.fill = green if cell.value == "Yes" else red
                cell.alignment = Alignment(horizontal="center")
            sheet.cell(row, len(headers)).fill = green if result == "PASS" else red
            sheet.cell(row, len(headers)).font = Font(bold=True)
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate([34, 24, 10, *([12] * 12), 16, 55, 12], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    coverage_sheet("Supplier Coverage", "Supplier", supplier)
    coverage_sheet("MR Coverage", "MR Source Country", mr)
    summary = workbook.create_sheet("Summary", 0)
    summary.append(["Dataset", "Input File", "Combinations", "Complete", "Incomplete", "Invalid Period Rows"])
    for cell in summary[1]:
        cell.fill, cell.font = blue, Font(color="FFFFFF", bold=True)
    for label, path, data, invalid in (("Supplier", supplier_input, supplier, supplier_invalid),
                                        ("Market Research", mr_input, mr, mr_invalid)):
        complete = sum(passed(months) for months in data.values())
        summary.append([label, str(path), len(data), complete, len(data) - complete, invalid])
    all_months = [*supplier.values(), *mr.values()]
    summary.append(["TOTAL", "", len(all_months), sum(passed(v) for v in all_months),
                    sum(not passed(v) for v in all_months), supplier_invalid + mr_invalid])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for index, width in enumerate([20, 95, 18, 14, 14, 20], 1):
        summary.column_dimensions[get_column_letter(index)].width = width
    summary.freeze_panes = "A2"
    workbook.save(output)


def print_result(label: str, coverage: Coverage, invalid: int) -> int:
    failures = [(key, EXPECTED - months) for key, months in coverage.items() if not passed(months)]
    print(f"{label}: checked={len(coverage)}, complete={len(coverage)-len(failures)}, incomplete={len(failures)}, invalid-period-rows={invalid}")
    for (entity, market, year), missing in sorted(failures):
        print(f"- {entity} | {market} | {year}: {', '.join(MONTHS[m-1] for m in sorted(missing))}")
    return len(failures)


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate months in Supplier and MR standardized models and export Excel.")
    parser.add_argument("--supplier-input", "--input", dest="supplier_input", type=Path, default=SUPPLIER_INPUT)
    parser.add_argument("--mr-input", type=Path, default=MR_INPUT)
    parser.add_argument("--year", type=int, help="Validate only this year")
    parser.add_argument("--supplier", help="Filter Supplier Name/source country in both models")
    parser.add_argument("--market", help="Filter Destination Country in both models")
    parser.add_argument("--output-excel", type=Path, help="Excel report path")
    args = parser.parse_args()
    output = args.output_excel or Path(f"month_coverage_validation_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    for label, path in (("Supplier", args.supplier_input), ("MR", args.mr_input)):
        if not path.is_file():
            print(f"ERROR: {label} input not found: {path}", file=sys.stderr)
            return 2
    try:
        supplier, supplier_invalid = validate(read_rows(args.supplier_input), args.year, args.supplier, args.market)
        mr, mr_invalid = validate(read_rows(args.mr_input), args.year, args.supplier, args.market)
        if not supplier and not mr:
            print("ERROR: No combinations matched the filters.", file=sys.stderr)
            return 2
        write_excel(output, supplier, mr, args.supplier_input, args.mr_input, supplier_invalid, mr_invalid)
    except (OSError, ValueError, ImportError, StopIteration) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    failures = print_result("Supplier", supplier, supplier_invalid)
    failures += print_result("Market Research", mr, mr_invalid)
    print(f"Excel report: {output.resolve()}")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
