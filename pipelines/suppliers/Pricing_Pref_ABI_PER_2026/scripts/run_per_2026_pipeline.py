"""
Pipeline: Peru 2026 (all months)
Reads each monthly file from data/raw/suppliers/PERU/2026/
Extracts from the Inputs sheet:
  - FOB Price (H9)
  - Ocean Freight (H10)
  - Import Clearance (H12 - Nacionalización)
  - DDP Price (H13 - Valor vPET)
Outputs standardized rows with Destination Country = "Peru".
"""
from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[4]
SOURCE_DIR = REPO_ROOT / "data" / "raw" / "suppliers" / "PERU" / "2026"
OUTPUT_XLSX = (
    REPO_ROOT / "data" / "processed" / "current" / "supplier"
    / "Data_Standardized_Front_End_Data_Model.xlsx"
)

SUPPLIER_NAME = "San Miguel Industrias (SMI)"
DESTINATION_COUNTRY = "Peru"
LOCATION = "Peru"
RESIN_INDEX_TYPE = "ICIS Asia SE Low index (M-2)"
TLC_FORMULA = "PET Resin USD/Ton = (ICIS+OF)*(1+IT)+CC"

MONTH_LABELS = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

COMPONENT_SPEC = [
    ("fob", "FOB Price", "Resin Index vPET", "Yes"),
    ("freight", "Ocean Freight", "Freight ", "Yes"),
    ("clearance", "Import Clearance (%)", "Tax", "Yes"),
    ("ddp", "DDP Price", "Total Landing Cost ", "Yes"),
]

FRONT_END_HEADERS = [
    "Data Type",
    "Source File ",
    "Supplier Name",
    "Destination Country",
    "Time_Period ",
    "Time Period Year",
    "Time Period Month",
    "Location",
    "Raw Cost Breakdown",
    "Resin Index Type",
    "Forecast Resin Index Type",
    "Mapping Columns",
    "Column Required for Calculation",
    "Value ",
    "TLC Formula",
]


def extract_month(path: Path) -> list[dict[str, Any]]:
    """Extract standardized rows from a single monthly Peru file."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Inputs"]

    year = int(ws.cell(6, 3).value)  # C6
    month = int(ws.cell(6, 4).value)  # D6
    period_label = f"{MONTH_LABELS[month]} {year}"

    fob = ws.cell(9, 8).value      # H9 - Precio FOB
    freight = ws.cell(10, 8).value  # H10 - Flete marítimo
    clearance = ws.cell(12, 8).value  # H12 - Nacionalización
    ddp = ws.cell(13, 8).value     # H13 - Valor vPET (DDP)

    values = {
        "fob": fob,
        "freight": freight,
        "clearance": clearance,
        "ddp": ddp,
    }

    rows: list[dict[str, Any]] = []
    for key, raw_label, mapping_col, required in COMPONENT_SPEC:
        rows.append({
            "Data Type": "Actual",
            "Source File ": path.name,
            "Supplier Name": SUPPLIER_NAME,
            "Destination Country": DESTINATION_COUNTRY,
            "Time_Period ": period_label,
            "Time Period Year": year,
            "Time Period Month": month,
            "Location": LOCATION,
            "Raw Cost Breakdown": raw_label,
            "Resin Index Type": RESIN_INDEX_TYPE,
            "Forecast Resin Index Type": "",
            "Mapping Columns": mapping_col,
            "Column Required for Calculation": required,
            "Value ": values[key],
            "TLC Formula": TLC_FORMULA,
        })

    wb.close()
    return rows


def main() -> None:
    if not SOURCE_DIR.exists():
        print(f"ERROR: Source directory not found: {SOURCE_DIR}")
        sys.exit(1)

    files = sorted(SOURCE_DIR.glob("*.xlsx"))
    if not files:
        print(f"ERROR: No xlsx files found in {SOURCE_DIR}")
        sys.exit(1)

    print(f"Found {len(files)} monthly files in {SOURCE_DIR.name}")

    all_rows: list[dict[str, Any]] = []
    for f in files:
        month_rows = extract_month(f)
        all_rows.extend(month_rows)
        year = month_rows[0]["Time Period Year"]
        month = month_rows[0]["Time Period Month"]
        print(f"  {f.name}: {MONTH_LABELS[month]} {year} -> {len(month_rows)} rows")

    print(f"\nTotal extracted rows: {len(all_rows)}")

    try:
        import pandas as pd
    except ImportError:
        print("ERROR: pandas required. Run: pip install pandas openpyxl")
        sys.exit(1)

    print(f"\nMerging into: {OUTPUT_XLSX}")
    existing_df = pd.read_excel(OUTPUT_XLSX, sheet_name="front_end_data_model")
    print(f"  Existing rows: {len(existing_df)}")

    # Remove old Peru rows
    old_count = len(existing_df[existing_df["Destination Country"] == DESTINATION_COUNTRY])
    if old_count > 0:
        existing_df = existing_df[existing_df["Destination Country"] != DESTINATION_COUNTRY]
        print(f"  Removed {old_count} old Peru rows")

    new_df = pd.DataFrame(all_rows, columns=FRONT_END_HEADERS)
    print(f"  New Peru rows: {len(new_df)}")

    merged_df = pd.concat([existing_df, new_df], ignore_index=True)
    print(f"  Merged total: {len(merged_df)} rows")

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        merged_df.to_excel(writer, sheet_name="front_end_data_model", index=False)
    print(f"  Written to: {OUTPUT_XLSX}")

    from collections import Counter
    dests = Counter(merged_df["Destination Country"].values)
    print(f"\nFinal destinations: {dict(dests)}")


if __name__ == "__main__":
    main()
