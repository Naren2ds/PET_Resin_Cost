from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

BASE = Path(__file__).resolve().parents[3]
SRC_XLSX = Path(r"C:/Users/C772391/Downloads/PET Resin Total landed cost calculator.xlsx")
OUT_XLSX = (
    BASE
    / "data"
    / "processed"
    / "current"
    / "supplier"
    / "All_Suppliers_Destinations_Taxes_Duties_Derivation.xlsx"
)

SUM_RANGE_RE = re.compile(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)(.*)", re.IGNORECASE)
CELL_REF_RE = re.compile(r"\b([A-Z]{1,3})(\d+)\b")


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize(value: str) -> str:
    return clean(value).lower()


def is_country_row(ws: object, row_num: int) -> bool:
    col_a = clean(ws.cell(row_num, 1).value)
    if not col_a:
        return False

    excluded = {
        "customs value",
        "total import taxes",
        "customs value: resin cost +freight+insurance",
    }
    if normalize(col_a) in excluded:
        return False

    return all(clean(ws.cell(row_num, c).value) == "" for c in range(2, 8))


def map_label(raw_label: str) -> str:
    label = normalize(raw_label)
    if label == "customs value":
        return ""

    if any(
        token in label
        for token in [
            "custom fee",
            "custom service",
            "customs insurance",
            "seguro",
            "fodinfa",
            "consular",
        ]
    ):
        return "customs clearance"

    if any(
        token in label
        for token in [
            "total import taxes",
            "tax",
            "vat",
            "iva",
            "igv",
            "cofins",
            "pis",
            "ipi",
            "duty",
            "impuesto",
            "ganancias",
            "ibb",
            "irae",
            "percepcion",
            "nacionalizacion",
        ]
    ):
        return "tax"

    return ""


def constant_or_variable(value: object) -> str:
    raw = clean(value)
    if raw.startswith("=") and re.search(r"[A-Z]+\d+", raw):
        return "Variable"

    try:
        float(str(value))
        return "Constant"
    except Exception:
        return "Variable" if raw else "Constant"


def prettify_formula_with_headers(
    formula: str,
    header_by_col: dict[str, str],
    value_row: int,
) -> str:
    expr = clean(formula)
    if not expr:
        return ""
    if expr.startswith("="):
        expr = expr[1:]

    expr = expr.replace("$", "")
    expr = expr.replace("'Calculation taxable base'!", "")
    expr = expr.replace("Calculation taxable base!", "")

    sum_match = SUM_RANGE_RE.fullmatch(expr)
    if sum_match:
        left_ref, right_ref, suffix = sum_match.groups()
        expr = f"({left_ref}+{right_ref}){suffix}"

    def replace_ref(match: re.Match[str]) -> str:
        col = match.group(1)
        row = int(match.group(2))
        if row == value_row and col in header_by_col and header_by_col[col]:
            return f"[{header_by_col[col]}]"
        return f"{col}{row}"

    return CELL_REF_RE.sub(replace_ref, expr)


def build_tax_duty_derivation() -> pd.DataFrame:
    wb = load_workbook(SRC_XLSX, data_only=False, read_only=True)
    ws = wb["Calculation taxable base"]

    rows: list[dict[str, str]] = []
    row_num = 1
    while row_num <= ws.max_row:
        if not is_country_row(ws, row_num):
            row_num += 1
            continue

        country = clean(ws.cell(row_num, 1).value)
        header_row = row_num + 1
        value_row = header_row + 1

        if value_row > ws.max_row or normalize(clean(ws.cell(header_row, 1).value)) != "customs value":
            row_num += 1
            continue

        header_by_col = {
            get_column_letter(c): clean(ws.cell(header_row, c).value)
            for c in range(1, 8)
        }

        for c in range(2, 8):
            raw_label = clean(ws.cell(header_row, c).value)
            if not raw_label:
                continue

            mapped = map_label(raw_label)
            if not mapped:
                continue

            cell_value = ws.cell(value_row, c).value
            if isinstance(cell_value, str):
                derived_formula = prettify_formula_with_headers(
                    formula=cell_value,
                    header_by_col=header_by_col,
                    value_row=value_row,
                )
            else:
                derived_formula = clean(cell_value)

            rows.append(
                {
                    "SupplierName": "PET Resin TLC Calculator",
                    "Destination COuntry": country,
                    "Raw Taxes & Duties(from supplier sheet)": raw_label,
                    "Mapped Taxes & Duties (what we have mapped )": mapped,
                    "Constant or Variable": constant_or_variable(cell_value),
                    "Derived Formula": derived_formula,
                }
            )

        # Include all "Total Import Taxes" variants for the current country block.
        totals_row = value_row + 1
        while totals_row <= ws.max_row and "total import taxes" in normalize(clean(ws.cell(totals_row, 1).value)):
            for label_col, formula_col in ((1, 2), (3, 4), (5, 6)):
                total_label = clean(ws.cell(totals_row, label_col).value)
                if not total_label:
                    continue
                if "total import taxes" not in normalize(total_label):
                    continue

                mapped = map_label(total_label) or "tax"
                formula_value = ws.cell(totals_row, formula_col).value
                if isinstance(formula_value, str):
                    derived_formula = prettify_formula_with_headers(
                        formula=formula_value,
                        header_by_col=header_by_col,
                        value_row=value_row,
                    )
                else:
                    derived_formula = clean(formula_value)

                rows.append(
                    {
                        "SupplierName": "PET Resin TLC Calculator",
                        "Destination COuntry": country,
                        "Raw Taxes & Duties(from supplier sheet)": total_label,
                        "Mapped Taxes & Duties (what we have mapped )": mapped,
                        "Constant or Variable": constant_or_variable(formula_value),
                        "Derived Formula": derived_formula,
                    }
                )

            totals_row += 1

        row_num = totals_row

    wb.close()

    out_df = pd.DataFrame(rows).sort_values(
        by=[
            "Destination COuntry",
            "Mapped Taxes & Duties (what we have mapped )",
            "Raw Taxes & Duties(from supplier sheet)",
        ]
    )

    return force_add_totals_for_core_countries(out_df)


def force_add_totals_for_core_countries(df: pd.DataFrame) -> pd.DataFrame:
    required_countries = {"brazil", "argentina", "uruguay"}
    vat_tokens = ("vat", "iva", "icms", "igv")

    def norm(value: str) -> str:
        return clean(value).lower()

    new_rows: list[dict[str, str]] = []

    for country_name, country_df in df.groupby("Destination COuntry", dropna=False):
        country_norm = norm(str(country_name))
        if country_norm not in required_countries:
            continue

        has_totals = country_df[
            country_df["Raw Taxes & Duties(from supplier sheet)"]
            .astype(str)
            .str.contains("Total Import Taxes", case=False, na=False)
        ]
        if not has_totals.empty:
            continue

        component_df = country_df[
            ~country_df["Raw Taxes & Duties(from supplier sheet)"]
            .astype(str)
            .str.contains("Total Import Taxes", case=False, na=False)
        ]

        if component_df.empty:
            continue

        labels: list[str] = []
        for label in component_df["Raw Taxes & Duties(from supplier sheet)"].astype(str).tolist():
            label_clean = clean(label)
            if label_clean and label_clean not in labels:
                labels.append(label_clean)

        if not labels:
            continue

        total_formula = f"({'+'.join(f'[{label}]' for label in labels)})/[Customs Value]"

        template = component_df.iloc[0].to_dict()
        template["Raw Taxes & Duties(from supplier sheet)"] = "Total Import Taxes"
        template["Mapped Taxes & Duties (what we have mapped )"] = "tax"
        template["Constant or Variable"] = "Variable"
        template["Derived Formula"] = total_formula
        new_rows.append(template)

        labels_excl_vat = [
            label for label in labels if not any(token in norm(label) for token in vat_tokens)
        ]
        if labels_excl_vat and len(labels_excl_vat) < len(labels):
            excl_formula = f"({'+'.join(f'[{label}]' for label in labels_excl_vat)})/[Customs Value]"
            template_excl = component_df.iloc[0].to_dict()
            template_excl["Raw Taxes & Duties(from supplier sheet)"] = "Total Import Taxes (Excl. VAT)"
            template_excl["Mapped Taxes & Duties (what we have mapped )"] = "tax"
            template_excl["Constant or Variable"] = "Variable"
            template_excl["Derived Formula"] = excl_formula
            new_rows.append(template_excl)

    if new_rows:
        df = pd.concat([df, pd.DataFrame(new_rows)], ignore_index=True)

    return df.sort_values(
        by=[
            "Destination COuntry",
            "Mapped Taxes & Duties (what we have mapped )",
            "Raw Taxes & Duties(from supplier sheet)",
        ]
    )


def main() -> None:
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    out_df = build_tax_duty_derivation()

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        out_df.to_excel(writer, sheet_name="tax_duty_mapping", index=False)

    print(f"Created: {OUT_XLSX}")
    print(f"Rows exported: {len(out_df)}")
    print(
        out_df[
            [
                "Destination COuntry",
                "Raw Taxes & Duties(from supplier sheet)",
                "Derived Formula",
            ]
        ]
        .head(20)
        .to_string(index=False)
    )


if __name__ == "__main__":
    main()
