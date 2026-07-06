from __future__ import annotations

import argparse
import csv
import json
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover
    load_dotenv = None

try:
    from openai import OpenAI
except Exception:  # pragma: no cover
    OpenAI = None


CANONICAL_HEADERS = [
    "Data Type",
    "Source File ",
    "Supplier Name",
    "Destination Country",
    "Time_Period ",
    "Time Period Year",
    "Time Period Month",
    "Location",
    "Raw Cost Breakdown",
    "Resin Index Type",
    "Forecast Resin Index Type",
    "Mapping Columns",
    "Column Required for Calculation",
    "Value ",
    "TLC Formula",
]

HEADER_ALIASES = {
    "period": {"period", "time period", "month", "mes", "fecha", "competencia"},
    "metric": {"metric", "component", "raw cost breakdown", "item", "concepto", "rubro"},
    "value": {"value", "valor", "usd", "amount", "importe", "precio"},
}

RAW_COST_SYNONYMS = {
    "resin": "PET resin cost (FOB)",
    "pet": "PET resin cost (FOB)",
    "ihs": "PET resin cost (FOB)",
    "icis": "PET resin cost (FOB)",
    "fob": "PET resin cost (FOB)",
    "precio fob": "PET resin cost (FOB)",
    "fob price": "PET resin cost (FOB)",
    "freight": "International freight",
    "flete": "International freight",
    "ocean freight": "International freight",
    "flete maritimo": "International freight",
    "frete internacional": "International freight",
    "tax": "Tax and Duties",
    "duty": "Tax and Duties",
    "impuesto": "Tax and Duties",
    "clearance": "Custom clearance",
    "nacionalizacion": "Custom clearance",
    "import clearance": "Custom clearance",
    "cif": "CIF price",
    "scrap": "Scrap",
    "others": "Others",
    "other": "Others",
    "seguro": "Insurance",
    "insurance": "Insurance",
    "tlc": "Sell Side Resina USD",
    "vpet": "Sell Side Resina USD",
    "valor vpet": "Sell Side Resina USD",
    "precio ddp": "Sell Side Resina USD",
    "total landing cost": "Sell Side Resina USD",
    "total resin price abi virgin formula": "Sell Side Resina USD",
}

INPUT_LABEL_TOKENS = {
    "precio fob",
    "fob price",
    "flete maritimo",
    "ocean freight",
    "precio cif",
    "cif price",
    "nacionalizacion",
    "import clearance",
    "valor vpet",
    "vpet value",
    "precio ddp",
    "ddp price",
}

INPUT_LABEL_EXCLUDE_TOKENS = {
    "merma",
    "discount",
    "descuento",
}

MONTH_ALIASES = {
    "jan": 1,
    "january": 1,
    "ene": 1,
    "enero": 1,
    "feb": 2,
    "february": 2,
    "febrero": 2,
    "mar": 3,
    "march": 3,
    "marzo": 3,
    "apr": 4,
    "april": 4,
    "abr": 4,
    "abril": 4,
    "may": 5,
    "mayo": 5,
    "jun": 6,
    "june": 6,
    "junio": 6,
    "jul": 7,
    "july": 7,
    "julio": 7,
    "aug": 8,
    "august": 8,
    "ago": 8,
    "agosto": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "set": 9,
    "septiembre": 9,
    "oct": 10,
    "october": 10,
    "octubre": 10,
    "nov": 11,
    "november": 11,
    "noviembre": 11,
    "dec": 12,
    "december": 12,
    "dic": 12,
    "diciembre": 12,
}


@dataclass
class ExtractedRow:
    sheet_name: str
    source_ref: str
    period_label: str
    year: int
    month: int
    raw_cost_breakdown: str
    value: float
    mapping_columns: str
    calc_column: str


@dataclass
class ValidationIssue:
    severity: str
    code: str
    message: str


class SupplierAsimovPipeline:
    def __init__(self, args: argparse.Namespace) -> None:
        self.args = args
        self.source_file = Path(args.source_file).resolve()
        self.output_dir = Path(args.output_dir).resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

        if load_dotenv is not None and args.dotenv_file:
            load_dotenv(dotenv_path=Path(args.dotenv_file).resolve())

    @staticmethod
    def _strip_accents(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value)
        return "".join(char for char in normalized if not unicodedata.combining(char))

    @staticmethod
    def _clean_text(value: Any) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()

    @staticmethod
    def _to_float(value: Any) -> float | None:
        if value in (None, ""):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(",", "")
        try:
            return float(text)
        except ValueError:
            return None

    def _period_from_value(self, value: Any) -> tuple[int, int, str] | None:
        if isinstance(value, datetime):
            period = date(value.year, value.month, 1)
            return period.year, period.month, period.strftime("%B %Y")
        if isinstance(value, date):
            period = date(value.year, value.month, 1)
            return period.year, period.month, period.strftime("%B %Y")

        text = self._clean_text(value)
        if not text:
            return None

        match = re.fullmatch(r"(\d{4})[-/](\d{1,2})", text)
        if match:
            year = int(match.group(1))
            month = int(match.group(2))
            if 1 <= month <= 12:
                return year, month, date(year, month, 1).strftime("%B %Y")

        token_match = re.search(r"([A-Za-z]{3,10})[\s\-_/]*(\d{2,4})", text)
        if token_match:
            month_token = token_match.group(1).lower()
            year = int(token_match.group(2))
            if year < 100:
                year += 2000
            month = MONTH_ALIASES.get(month_token)
            if month:
                return year, month, date(year, month, 1).strftime("%B %Y")

        return None

    @staticmethod
    def _normalize_key(value: Any) -> str:
        text = SupplierAsimovPipeline._strip_accents(str(value or "")).lower()
        return re.sub(r"[^a-z0-9]+", " ", text).strip()

    def _map_raw_cost(self, metric: str) -> str:
        norm = self._normalize_key(metric)
        if norm in RAW_COST_SYNONYMS:
            return RAW_COST_SYNONYMS[norm]
        for token, canonical in RAW_COST_SYNONYMS.items():
            if token in norm:
                return canonical
        return ""

    def _detect_header_row(self, worksheet: Any, max_scan_rows: int = 20) -> tuple[int, dict[str, int]] | None:
        for row_idx in range(1, min(max_scan_rows, worksheet.max_row) + 1):
            header_map: dict[str, int] = {}
            cells = [self._clean_text(worksheet.cell(row_idx, col).value).lower() for col in range(1, worksheet.max_column + 1)]
            for col_idx, raw_header in enumerate(cells, start=1):
                norm = self._normalize_key(raw_header)
                for target, aliases in HEADER_ALIASES.items():
                    if norm in aliases and target not in header_map:
                        header_map[target] = col_idx
            if all(field in header_map for field in ("period", "metric", "value")):
                return row_idx, header_map
        return None

    def _profile_workbook(self) -> dict[str, Any]:
        workbook = load_workbook(self.source_file, data_only=False, read_only=True)
        try:
            sheets: list[dict[str, Any]] = []
            for ws in workbook.worksheets:
                axis = self._detect_month_axis(ws)
                sheets.append(
                    {
                        "sheet_name": ws.title,
                        "max_row": ws.max_row,
                        "max_column": ws.max_column,
                        "header_detected": bool(self._detect_header_row(ws)),
                        "month_axis_detected": bool(axis),
                        "month_axis_row": axis[0] if axis else None,
                        "month_axis_columns": len(axis[1]) if axis else 0,
                    }
                )
            return {
                "source_file": str(self.source_file),
                "sheet_count": len(workbook.sheetnames),
                "sheets": sheets,
            }
        finally:
            workbook.close()

    def _deterministic_extract(self) -> tuple[list[ExtractedRow], list[dict[str, Any]]]:
        workbook = load_workbook(self.source_file, data_only=True, read_only=True)
        extracted: list[ExtractedRow] = []
        unresolved: list[dict[str, Any]] = []
        try:
            for ws in workbook.worksheets:
                detected = self._detect_header_row(ws)
                if not detected:
                    continue
                header_row, header_map = detected

                for row_idx in range(header_row + 1, ws.max_row + 1):
                    period_raw = ws.cell(row_idx, header_map["period"]).value
                    metric_raw = self._clean_text(ws.cell(row_idx, header_map["metric"]).value)
                    value_raw = ws.cell(row_idx, header_map["value"]).value

                    period = self._period_from_value(period_raw)
                    amount = self._to_float(value_raw)
                    mapped_metric = self._map_raw_cost(metric_raw)

                    if period and amount is not None and mapped_metric:
                        year, month, label = period
                        extracted.append(
                            ExtractedRow(
                                sheet_name=ws.title,
                                source_ref=f"{ws.title}!R{row_idx}",
                                period_label=label,
                                year=year,
                                month=month,
                                raw_cost_breakdown=mapped_metric,
                                value=amount,
                                mapping_columns=metric_raw,
                                calc_column=metric_raw,
                            )
                        )
                    else:
                        unresolved.append(
                            {
                                "sheet_name": ws.title,
                                "row": row_idx,
                                "period_raw": self._clean_text(period_raw),
                                "metric_raw": metric_raw,
                                "value_raw": self._clean_text(value_raw),
                                "period_ok": bool(period),
                                "value_ok": amount is not None,
                                "mapped_metric": mapped_metric,
                            }
                        )
        finally:
            workbook.close()

        return extracted, unresolved

    def _detect_month_axis(self, worksheet: Any) -> tuple[int, dict[int, tuple[int, int, str]]] | None:
        scan_limit = min(70, worksheet.max_row)
        for row_idx in range(1, scan_limit + 1):
            periods_by_col: dict[int, tuple[int, int, str]] = {}
            for col_idx in range(1, worksheet.max_column + 1):
                parsed = self._period_from_value(worksheet.cell(row_idx, col_idx).value)
                if parsed:
                    periods_by_col[col_idx] = parsed
            if len(periods_by_col) >= 4:
                return row_idx, periods_by_col
        return None

    def _grid_extract(self) -> list[ExtractedRow]:
        workbook = load_workbook(self.source_file, data_only=True, read_only=True)
        extracted: list[ExtractedRow] = []
        try:
            for ws in workbook.worksheets:
                axis = self._detect_month_axis(ws)
                if not axis:
                    continue

                axis_row, periods_by_col = axis
                max_row = min(ws.max_row, axis_row + 140)

                for row_idx in range(axis_row + 1, max_row + 1):
                    label = ""
                    for metric_col in (1, 2, 3, 4):
                        label = self._clean_text(ws.cell(row_idx, metric_col).value)
                        if label:
                            break

                    if not label:
                        continue

                    mapped_metric = self._map_raw_cost(label)
                    if not mapped_metric:
                        continue

                    for col_idx, (year, month, period_label) in periods_by_col.items():
                        value = self._to_float(ws.cell(row_idx, col_idx).value)
                        if value is None:
                            continue
                        extracted.append(
                            ExtractedRow(
                                sheet_name=ws.title,
                                source_ref=f"{ws.title}!R{row_idx}C{col_idx}",
                                period_label=period_label,
                                year=year,
                                month=month,
                                raw_cost_breakdown=mapped_metric,
                                value=value,
                                mapping_columns=label,
                                calc_column=label,
                            )
                        )
        finally:
            workbook.close()

        return extracted

    def _workbook_period_hint(self, worksheet: Any) -> tuple[int, int, str] | None:
        # Pricing_Pref style workbooks keep year/month in C6/D6.
        period = self._period_from_value(f"{worksheet['C6'].value}-{worksheet['D6'].value}")
        if period:
            return period

        for candidate in (worksheet["C9"].value, worksheet["C7"].value, worksheet["D6"].value):
            period = self._period_from_value(candidate)
            if period:
                return period
        return None

    def _is_inputs_sheet(self, sheet_name: str) -> bool:
        normalized = self._normalize_key(sheet_name)
        return normalized == "inputs" or "input" in normalized

    def _inputs_sheet_extract(self) -> list[ExtractedRow]:
        workbook = load_workbook(self.source_file, data_only=True, read_only=True)
        extracted: list[ExtractedRow] = []
        try:
            for ws in workbook.worksheets:
                if not self._is_inputs_sheet(ws.title):
                    continue

                period = self._workbook_period_hint(ws)
                if not period:
                    continue
                year, month, period_label = period

                max_row = min(ws.max_row, 260)
                max_col = min(ws.max_column, 20)

                for row_idx in range(1, max_row + 1):
                    for col_idx in range(1, max_col):
                        label = self._clean_text(ws.cell(row_idx, col_idx).value)
                        if not label:
                            continue
                        label_key = self._normalize_key(label)
                        if any(token in label_key for token in INPUT_LABEL_EXCLUDE_TOKENS):
                            continue
                        if label_key not in INPUT_LABEL_TOKENS:
                            continue

                        value = self._to_float(ws.cell(row_idx, col_idx + 1).value)
                        if value is None:
                            continue

                        mapped_metric = self._map_raw_cost(label)
                        if not mapped_metric:
                            continue

                        extracted.append(
                            ExtractedRow(
                                sheet_name=ws.title,
                                source_ref=f"{ws.title}!R{row_idx}C{col_idx + 1}",
                                period_label=period_label,
                                year=year,
                                month=month,
                                raw_cost_breakdown=mapped_metric,
                                value=value,
                                mapping_columns=label,
                                calc_column=label,
                            )
                        )
        finally:
            workbook.close()

        return extracted

    def _dedupe_extracted_rows(self, rows: list[ExtractedRow]) -> list[ExtractedRow]:
        deduped: dict[tuple[int, int, str], ExtractedRow] = {}
        for row in rows:
            key = (row.year, row.month, row.raw_cost_breakdown)
            if key not in deduped:
                deduped[key] = row
        return sorted(deduped.values(), key=lambda r: (r.year, r.month, r.raw_cost_breakdown))

    @staticmethod
    def _env_value(*names: str) -> str | None:
        for name in names:
            value = os.getenv(name)
            if value and value.strip():
                return value.strip()
        return None

    def _asimov_enrich(self, unresolved: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if not unresolved:
            return []
        if OpenAI is None:
            return []

        api_key = self._env_value("OPENAI_API_KEY", "ASIMOV_API_KEY")
        if not api_key:
            return []

        base_url = self._env_value("BASE_URL", "ASIMOV_BASE_URL")
        model = self.args.asimov_model
        max_records = max(1, int(self.args.asimov_max_records))
        payload = unresolved[:max_records]

        prompt = {
            "task": "Map supplier rows to canonical cost components and parse periods.",
            "canonical_costs": sorted(set(RAW_COST_SYNONYMS.values())),
            "rows": payload,
            "response_format": {
                "items": [
                    {
                        "sheet_name": "string",
                        "row": "int",
                        "raw_cost_breakdown": "canonical string or empty",
                        "time_period_year": "int or null",
                        "time_period_month": "int or null",
                        "time_period_label": "Month YYYY or empty",
                        "confidence": "0.0-1.0",
                    }
                ]
            },
        }

        client = OpenAI(api_key=api_key, base_url=base_url)
        response = client.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": "You are a strict extraction assistant. Return valid JSON only.",
                },
                {
                    "role": "user",
                    "content": json.dumps(prompt, ensure_ascii=True),
                },
            ],
            temperature=0,
            max_tokens=1500,
        )

        content = response.choices[0].message.content or ""
        try:
            parsed = json.loads(content)
            items = parsed.get("items", []) if isinstance(parsed, dict) else []
            return items if isinstance(items, list) else []
        except json.JSONDecodeError:
            return []

    def _merge_asimov_suggestions(
        self,
        extracted: list[ExtractedRow],
        unresolved: list[dict[str, Any]],
        suggestions: list[dict[str, Any]],
    ) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        unresolved_lookup = {(item["sheet_name"], int(item["row"])): item for item in unresolved}

        for item in suggestions:
            try:
                key = (str(item.get("sheet_name", "")), int(item.get("row", 0)))
                unresolved_row = unresolved_lookup.get(key)
                if not unresolved_row:
                    continue

                confidence = float(item.get("confidence", 0))
                if confidence < self.args.confidence_threshold:
                    continue

                raw_cost = self._clean_text(item.get("raw_cost_breakdown"))
                year = item.get("time_period_year")
                month = item.get("time_period_month")
                label = self._clean_text(item.get("time_period_label"))

                if not raw_cost or year is None or month is None:
                    continue

                mapped = self._map_raw_cost(raw_cost) or raw_cost
                value = self._to_float(unresolved_row.get("value_raw"))
                if value is None:
                    continue

                extracted.append(
                    ExtractedRow(
                        sheet_name=key[0],
                        source_ref=f"{key[0]}!R{key[1]}",
                        period_label=label or date(int(year), int(month), 1).strftime("%B %Y"),
                        year=int(year),
                        month=int(month),
                        raw_cost_breakdown=mapped,
                        value=value,
                        mapping_columns=self._clean_text(unresolved_row.get("metric_raw")),
                        calc_column=self._clean_text(unresolved_row.get("metric_raw")),
                    )
                )
            except Exception as ex:  # pragma: no cover
                issues.append(
                    ValidationIssue(
                        severity="warning",
                        code="ASIMOV_MERGE_SKIPPED",
                        message=f"Failed to merge one Asimov suggestion: {ex}",
                    )
                )

        return issues

    def _to_front_end_rows(self, extracted: list[ExtractedRow]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        location = self.args.location or self.args.destination_country

        for row in extracted:
            rows.append(
                {
                    "Data Type": "actual",
                    "Source File ": self.source_file.name,
                    "Supplier Name": self.args.supplier_name,
                    "Destination Country": self.args.destination_country,
                    "Time_Period ": row.period_label,
                    "Time Period Year": row.year,
                    "Time Period Month": row.month,
                    "Location": location,
                    "Raw Cost Breakdown": row.raw_cost_breakdown,
                    "Resin Index Type": self.args.resin_index_type,
                    "Forecast Resin Index Type": self.args.forecast_resin_index_type,
                    "Mapping Columns": row.mapping_columns,
                    "Column Required for Calculation": row.calc_column,
                    "Value ": row.value,
                    "TLC Formula": self.args.tlc_formula,
                }
            )

        return rows

    def _validate_rows(self, rows: list[dict[str, Any]]) -> list[ValidationIssue]:
        issues: list[ValidationIssue] = []
        seen: set[tuple[Any, ...]] = set()

        required = [
            "Data Type",
            "Supplier Name",
            "Destination Country",
            "Time_Period ",
            "Time Period Year",
            "Time Period Month",
            "Raw Cost Breakdown",
            "Value ",
        ]

        for idx, row in enumerate(rows, start=2):
            for field in required:
                if self._clean_text(row.get(field)) == "":
                    issues.append(
                        ValidationIssue(
                            severity="error",
                            code="MISSING_REQUIRED",
                            message=f"Row {idx}: missing required field '{field}'.",
                        )
                    )

            value = self._to_float(row.get("Value "))
            if value is None:
                issues.append(
                    ValidationIssue(
                        severity="error",
                        code="VALUE_NOT_NUMERIC",
                        message=f"Row {idx}: Value is not numeric.",
                    )
                )

            key = (
                row.get("Supplier Name"),
                row.get("Destination Country"),
                row.get("Location"),
                row.get("Raw Cost Breakdown"),
                row.get("Time Period Year"),
                row.get("Time Period Month"),
                row.get("Data Type"),
            )
            if key in seen:
                issues.append(
                    ValidationIssue(
                        severity="error",
                        code="DUPLICATE_ROW",
                        message=f"Row {idx}: duplicate grain detected {key}.",
                    )
                )
            seen.add(key)

        return issues

    def _write_csv(self, path: Path, rows: list[dict[str, Any]]) -> None:
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=CANONICAL_HEADERS)
            writer.writeheader()
            for row in rows:
                writer.writerow({header: row.get(header, "") for header in CANONICAL_HEADERS})

    def _write_xlsx(self, path: Path, rows: list[dict[str, Any]]) -> None:
        workbook = Workbook()
        ws = workbook.active
        ws.title = "front_end_standardized"
        ws.append(CANONICAL_HEADERS)
        for row in rows:
            ws.append([row.get(header, "") for header in CANONICAL_HEADERS])
        workbook.save(path)

    def _write_json(self, path: Path, payload: dict[str, Any]) -> None:
        path.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    def run(self) -> int:
        if not self.source_file.exists():
            raise FileNotFoundError(f"Source file not found: {self.source_file}")

        profile = self._profile_workbook()
        extracted, unresolved = self._deterministic_extract()
        issues: list[ValidationIssue] = []

        if not extracted:
            extracted = self._inputs_sheet_extract()

        # Phase 2 fallback: matrix/grid extraction for month-column supplier workbooks.
        if not extracted:
            extracted = self._grid_extract()

        if self.args.asimov_mode != "off" and unresolved:
            suggestions = self._asimov_enrich(unresolved)
            issues.extend(self._merge_asimov_suggestions(extracted, unresolved, suggestions))
            unresolved = self._remaining_unresolved(unresolved, extracted)

        extracted = self._dedupe_extracted_rows(extracted)

        if self.args.asimov_mode == "required" and unresolved:
            issues.append(
                ValidationIssue(
                    severity="error",
                    code="ASIMOV_REQUIRED_UNRESOLVED",
                    message=f"Asimov required mode failed to resolve {len(unresolved)} rows.",
                )
            )

        front_end_rows = self._to_front_end_rows(extracted)
        issues.extend(self._validate_rows(front_end_rows))

        # Guardrail: zero-row extraction indicates unsupported workbook layout.
        if not front_end_rows:
            issues.append(
                ValidationIssue(
                    severity="error",
                    code="NO_ROWS_EXTRACTED",
                    message=(
                        "No rows were extracted from the workbook. "
                        "This supplier layout is not currently supported by the generic extractor."
                    ),
                )
            )

        csv_path = self.output_dir / "generic_supplier_actual_front_end_standardized.csv"
        xlsx_path = self.output_dir / "generic_supplier_actual_front_end_standardized.xlsx"
        summary_path = self.output_dir / "generic_supplier_validation_summary.json"
        profile_path = self.output_dir / "generic_supplier_workbook_profile.json"
        unresolved_path = self.output_dir / "generic_supplier_unresolved_rows.json"

        self._write_csv(csv_path, front_end_rows)
        self._write_xlsx(xlsx_path, front_end_rows)
        self._write_json(profile_path, profile)
        self._write_json(unresolved_path, {"count": len(unresolved), "rows": unresolved})

        error_count = sum(1 for issue in issues if issue.severity == "error")
        warning_count = sum(1 for issue in issues if issue.severity == "warning")

        summary = {
            "source_file": str(self.source_file),
            "output_csv": str(csv_path),
            "output_xlsx": str(xlsx_path),
            "profile_file": str(profile_path),
            "unresolved_file": str(unresolved_path),
            "supplier_name": self.args.supplier_name,
            "destination_country": self.args.destination_country,
            "asimov_mode": self.args.asimov_mode,
            "asimov_model": self.args.asimov_model,
            "rows_extracted": len(extracted),
            "rows_output": len(front_end_rows),
            "unresolved_rows": len(unresolved),
            "error_count": error_count,
            "warning_count": warning_count,
            "issues": [issue.__dict__ for issue in issues],
        }
        self._write_json(summary_path, summary)

        print(json.dumps(summary, indent=2))
        return 0 if error_count == 0 else 2

    def _remaining_unresolved(
        self,
        unresolved: list[dict[str, Any]],
        extracted: list[ExtractedRow],
    ) -> list[dict[str, Any]]:
        resolved_keys = {(row.sheet_name, int(row.source_ref.split("!R")[-1])) for row in extracted}
        remaining: list[dict[str, Any]] = []
        for item in unresolved:
            key = (str(item.get("sheet_name", "")), int(item.get("row", 0)))
            if key not in resolved_keys:
                remaining.append(item)
        return remaining


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Generic supplier extraction pipeline (single file) with deterministic extraction "
            "plus optional Asimov-assisted fallback for unresolved rows."
        )
    )
    parser.add_argument("--source-file", required=True, help="Supplier source workbook path.")
    parser.add_argument(
        "--output-dir",
        default=str(Path("pipelines/suppliers/generic_asimov_single_file/artifacts")),
        help="Output directory for standardized files and validation outputs.",
    )
    parser.add_argument("--supplier-name", required=True, help="Supplier name for standardized output.")
    parser.add_argument(
        "--destination-country",
        required=True,
        help="Destination country for standardized output.",
    )
    parser.add_argument(
        "--location",
        default="",
        help="Optional location. Defaults to destination country.",
    )
    parser.add_argument(
        "--resin-index-type",
        default="",
        help="Resin index type to stamp in output rows.",
    )
    parser.add_argument(
        "--forecast-resin-index-type",
        default="",
        help="Forecast resin index type column value.",
    )
    parser.add_argument(
        "--tlc-formula",
        default="",
        help="TLC formula metadata to stamp in standardized rows.",
    )
    parser.add_argument(
        "--asimov-mode",
        choices=["off", "assist", "required"],
        default="assist",
        help="off=deterministic only, assist=use Asimov for unresolved rows, required=fail if unresolved rows remain.",
    )
    parser.add_argument(
        "--asimov-model",
        default="openai/gpt-4o",
        help="Asimov/OpenAI-compatible model identifier.",
    )
    parser.add_argument(
        "--asimov-max-records",
        type=int,
        default=150,
        help="Max unresolved rows sent to Asimov in one call.",
    )
    parser.add_argument(
        "--confidence-threshold",
        type=float,
        default=0.80,
        help="Minimum Asimov confidence to accept row suggestion.",
    )
    parser.add_argument(
        "--dotenv-file",
        default="apps/backend/.env",
        help="Dotenv file for OPENAI_API_KEY/ASIMOV_API_KEY and BASE_URL/ASIMOV_BASE_URL.",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    pipeline = SupplierAsimovPipeline(args)
    return pipeline.run()


if __name__ == "__main__":
    raise SystemExit(main())
