from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[4]
OPPORTUNITY_WORKBOOK = (
    REPO_ROOT
    / "data"
    / "processed"
    / "current"
    / "opportunity"
    / "PET_Quantification_Breakdown_By_Destination.xlsx"
)

SUMMARY_SHEET = "Summary_Opportunity"
NON_DESTINATION_SHEETS = {"Actuals_Reference", SUMMARY_SHEET}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\ufffd", "–").strip()


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _summary_record(row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        "destination": _text(row[0]),
        "supplierLocation": _text(row[1]),
        "volume": _number(row[2]),
        "supplierPeriod": _text(row[3]),
        "supplierDataType": _text(row[4]),
        "supplierTlc": _number(row[5]),
        "alternativeSource": _text(row[6]),
        "marketPeriod": _text(row[7]),
        "marketDataType": _text(row[8]),
        "marketTlc": _number(row[9]),
        "marketMinusSupplierDelta": _number(row[10]),
        "currentSpend": _number(row[11]),
        "alternativeSpend": _number(row[12]),
        "estimatedOpportunity": _number(row[13]),
        "status": _text(row[14]),
        "routeSelectionRule": _text(row[15]),
        "primaryCostDriver": _text(row[16]),
        "supplierDriverCost": _number(row[17]),
        "marketDriverCost": _number(row[18]),
        "driverGap": _number(row[19]),
        "driverInterpretation": _text(row[20]),
    }


def _breakdown_record(row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        "destination": _text(row[0]),
        "supplierLocation": _text(row[1]),
        "volume": _number(row[2]),
        "alternativeSource": _text(row[3]),
        "component": _text(row[4]),
        "marketCost": _number(row[5]),
        "supplierCost": _number(row[6]),
        "marketMinusSupplierDifference": _number(row[7]),
        "calculation": _text(row[8]),
        "marketImpact": _number(row[9]),
        "savingsContribution": _number(row[10]),
        "interpretation": _text(row[11]),
    }


def _amcor_scenario(entries: list[dict[str, Any]], location: str) -> dict[str, Any] | None:
    location_key = location.casefold()
    return next(
        (
            entry
            for entry in entries
            if entry["destination"] == "Brazil"
            and "amcor" in entry["supplierLocation"].casefold()
            and location_key in entry["supplierLocation"].casefold()
        ),
        None,
    )


def _metrics(entries: list[dict[str, Any]]) -> dict[str, Any]:
    destinations = sorted({entry["destination"] for entry in entries})
    non_brazil = [entry for entry in entries if entry["destination"] != "Brazil"]
    brazil_non_amcor = [
        entry
        for entry in entries
        if entry["destination"] == "Brazil"
        and "amcor" not in entry["supplierLocation"].casefold()
    ]
    amcor_manaus = _amcor_scenario(entries, "manaus")
    amcor_suape = _amcor_scenario(entries, "suape")

    base_opportunity = sum(
        entry["estimatedOpportunity"] for entry in non_brazil + brazil_non_amcor
    )
    amcor_opportunities = [
        entry["estimatedOpportunity"]
        for entry in (amcor_manaus, amcor_suape)
        if entry is not None
    ]
    opportunity_low = base_opportunity + (min(amcor_opportunities) if amcor_opportunities else 0)
    opportunity_high = base_opportunity + (max(amcor_opportunities) if amcor_opportunities else 0)

    base_volume = sum(entry["volume"] for entry in non_brazil + brazil_non_amcor)
    amcor_volumes = [
        entry["volume"]
        for entry in (amcor_manaus, amcor_suape)
        if entry is not None
    ]
    assessed_volume = base_volume + (max(amcor_volumes) if amcor_volumes else 0)

    positive_destinations = len(
        {
            entry["destination"]
            for entry in entries
            if entry["estimatedOpportunity"] > 0
        }
    )

    periods = sorted(
        {
            entry["supplierPeriod"]
            for entry in entries
            if entry["supplierPeriod"]
        }
    )
    return {
        "destinationCount": len(destinations),
        "positiveDestinationCount": positive_destinations,
        "assessedVolume": assessed_volume,
        "opportunityLow": opportunity_low,
        "opportunityHigh": opportunity_high,
        "period": periods[-1] if periods else "",
        "brazilScenarioNote": (
            "Brazil Amcor uses one shared 16,912 MT volume. Manaus and Suape are "
            "presented as alternative scenarios and are not added together."
        ),
    }


@lru_cache(maxsize=2)
def _read_opportunity_data(_modified_time: float) -> dict[str, Any]:
    workbook = load_workbook(OPPORTUNITY_WORKBOOK, read_only=True, data_only=True)
    try:
        summary_sheet = workbook[SUMMARY_SHEET]
        summary = [
            _summary_record(row)
            for row in summary_sheet.iter_rows(min_row=2, values_only=True)
            if _text(row[0])
        ]

        breakdowns: list[dict[str, Any]] = []
        for sheet_name in workbook.sheetnames:
            if sheet_name in NON_DESTINATION_SHEETS:
                continue
            sheet = workbook[sheet_name]
            breakdowns.extend(
                _breakdown_record(row)
                for row in sheet.iter_rows(min_row=3, values_only=True)
                if _text(row[0]) == sheet_name and _text(row[4])
            )
    finally:
        workbook.close()

    summary.sort(key=lambda entry: entry["estimatedOpportunity"], reverse=True)
    return {
        "sourceFile": str(OPPORTUNITY_WORKBOOK.relative_to(REPO_ROOT)),
        "metrics": _metrics(summary),
        "summary": summary,
        "breakdowns": breakdowns,
    }


def opportunity_data() -> dict[str, Any]:
    if not OPPORTUNITY_WORKBOOK.exists():
        raise FileNotFoundError(OPPORTUNITY_WORKBOOK)
    return _read_opportunity_data(OPPORTUNITY_WORKBOOK.stat().st_mtime)


def opportunity_workbook_path() -> Path:
    if not OPPORTUNITY_WORKBOOK.exists():
        raise FileNotFoundError(OPPORTUNITY_WORKBOOK)
    return OPPORTUNITY_WORKBOOK
