"""
Pipeline: El Salvador and Honduras 2026 (CORRECTED)
Source: Sell_Side_2026-ABI May 2026 AMCOR (1).xlsx
Sheet: ICIS Price History
Uses column S for month/year reference (not column B which has weekly ICIS dates).
Extracts rows where S has a 2026 date.

Components (from mapping file):
  - I: PET Bottle Grade FOB China Mid (N-4) -> Resin Index vPET
  - K: Adder -> Tax
  - L: Logistics -> Freight
  - M: LANDED -> Tax
  - N: Finance fee V (SOFR +3,8%@185) -> Insurance
  - R: Total precio ABI -> Total Landing Cost

Destination: El Salvador and Honduras
Supplier: Amcor
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[4]
SOURCE_FILE = REPO_ROOT / "data" / "raw" / "suppliers" / "Sell_Side_2026-ABI May 2026 AMCOR (1).xlsx"
OUTPUT_XLSX = (
    REPO_ROOT / "data" / "processed" / "current" / "supplier"
    / "Data_Standardized_Front_End_Data_Model.xlsx"
)

SUPPLIER_NAME = "Amcor"
DESTINATION_COUNTRY = "El Salvador and Honduras"
LOCATION = "China"
RESIN_INDEX_TYPE = "PET Bottle Grade FOB China Mid (N-4) USD/ton"
TLC_FORMULA = "Total precio ABI = Index(N-4) + Adder + Logistics + LANDED + Finance fee"

MONTH_LABELS = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

# Column positions (1-indexed)
COL_S_MONTH = 19    # S: Month/Year reference
COL_INDEX = 9       # I: ICIS Index
COL_ADDER = 11      # K: Adder
COL_LOGISTICS = 12  # L: Logistics
COL_LANDED = 13     # M: LANDED
COL_FINANCE = 14    # N: Finance fee
COL_TOTAL = 18      # R: Total precio ABI

# Components to extract: (key, raw_label, col_idx, mapping_col, required)
COMPONENT_SPEC = [
    ("index", "PET Bottle Grade FOB China Mid (N-4)", COL_INDEX, "Resin Index vPET", "Yes"),
    ("adder", "Adder", COL_ADDER, "Tax", "Yes"),
    ("logistics", "Logistics", COL_LOGISTICS, "Freight ", "Yes"),
    ("landed", "LANDED", COL_LANDED, "Tax", "Yes"),
    ("finance", "Finance fee V (SOFR +3,8%@185)", COL_FINANCE, "Insurance", "Yes"),
    ("total", "Total precio ABI", COL_TOTAL, "Total Landing Cost ", "Yes"),
]

FRONT_END_HEADERS = [
    "Data Type",
    "Source File ",
    "Supplier Name",
    "Destination Country",
    "Time_Period ",
    "Time Period Year",
    "Time Period Month",
    "Location",
    "Raw Cost Breakdown",
    "Resin Index Type",
    "Forecast Resin Index Type",
    "Mapping Columns",
    "Column Required for Calculation",
    "Value ",
    "TLC Formula",
]


def parse_date(val) -> tuple[int, int] | None:
    """Parse date value from column S, return (year, month)."""
    if isinstance(val, datetime):
        return val.year, val.month
    if hasattr(val, 'year'):
        return val.year, val.month
    if isinstance(val, str):
        val = val.strip()
        if not val:
            return None
        try:
            dt = datetime.fromisoformat(val.replace(" 00:00:00", ""))
            return dt.year, dt.month
        except ValueError:
            pass
    return None


def extract_2026_data() -> list[dict[str, Any]]:
    """Extract rows from ICIS Price History where column S has a 2026 date."""
    wb = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    ws = wb["ICIS Price History"]

    all_rows: list[dict[str, Any]] = []

    for row_idx in range(14, 300):
        s_val = ws.cell(row_idx, COL_S_MONTH).value
        if s_val is None:
            continue

        parsed = parse_date(s_val)
        if parsed is None:
            continue

        year, month = parsed
        if year != 2026:
            continue

        period_label = f"{MONTH_LABELS[month]} {year}"

        # Extract all component values
        values = {}
        for key, raw_label, col_idx, mapping_col, required in COMPONENT_SPEC:
            cell_val = ws.cell(row_idx, col_idx).value
            if cell_val is not None and isinstance(cell_val, (int, float)):
                values[key] = cell_val
            else:
                values[key] = None

        total = values.get("total")
        if total is None:
            continue

        print(f"  Row {row_idx}: {period_label} -> ICIS={values['index']}, Adder={values['adder']}, "
              f"Logistics={values['logistics']}, Landed={values['landed']}, "
              f"Finance={str(values['finance'])[:8]}, Total={total}")

        for key, raw_label, col_idx, mapping_col, required in COMPONENT_SPEC:
            val = values[key]
            if val is None:
                continue
            all_rows.append({
                "Data Type": "Actual",
                "Source File ": SOURCE_FILE.name,
                "Supplier Name": SUPPLIER_NAME,
                "Destination Country": DESTINATION_COUNTRY,
                "Time_Period ": period_label,
                "Time Period Year": year,
                "Time Period Month": month,
                "Location": LOCATION,
                "Raw Cost Breakdown": raw_label,
                "Resin Index Type": RESIN_INDEX_TYPE,
                "Forecast Resin Index Type": "",
                "Mapping Columns": mapping_col,
                "Column Required for Calculation": required,
                "Value ": round(val, 6) if isinstance(val, float) else val,
                "TLC Formula": TLC_FORMULA,
            })

    wb.close()
    return all_rows


def main() -> None:
    if not SOURCE_FILE.exists():
        print(f"ERROR: Source file not found: {SOURCE_FILE}")
        sys.exit(1)

    print(f"Reading: {SOURCE_FILE.name}")
    print(f"Sheet: ICIS Price History, Column S for month/year (2026 only)")
    print()

    all_rows = extract_2026_data()
    print(f"\nTotal extracted rows: {len(all_rows)}")

    if not all_rows:
        print("ERROR: No rows extracted")
        sys.exit(1)

    try:
        import pandas as pd
    except ImportError:
        print("ERROR: pandas required. Run: pip install pandas openpyxl")
        sys.exit(1)

    print(f"\nMerging into: {OUTPUT_XLSX}")
    existing_df = pd.read_excel(OUTPUT_XLSX, sheet_name="front_end_data_model")
    print(f"  Existing rows: {len(existing_df)}")

    # Remove old El Salvador and Honduras rows
    old_count = len(existing_df[existing_df["Destination Country"] == DESTINATION_COUNTRY])
    if old_count > 0:
        existing_df = existing_df[existing_df["Destination Country"] != DESTINATION_COUNTRY]
        print(f"  Removed {old_count} old El Salvador and Honduras rows")

    new_df = pd.DataFrame(all_rows, columns=FRONT_END_HEADERS)
    print(f"  New El Salvador and Honduras rows: {len(new_df)}")

    merged_df = pd.concat([existing_df, new_df], ignore_index=True)
    print(f"  Merged total: {len(merged_df)} rows")

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        merged_df.to_excel(writer, sheet_name="front_end_data_model", index=False)
    print(f"  Written to: {OUTPUT_XLSX}")

    from collections import Counter
    dests = Counter(merged_df["Destination Country"].values)
    print(f"\nFinal destinations: {dict(dests)}")


if __name__ == "__main__":
    main()
