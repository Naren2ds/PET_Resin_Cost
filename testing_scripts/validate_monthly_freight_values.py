"""Validate monthly Supplier and Market Research freight values.

Run from PET_Resin_Cost:
    python testing_scripts/validate_monthly_freight_values.py --year 2026
"""
from __future__ import annotations
import argparse, math, sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
SUPPLIER_INPUT = ROOT / "data/processed/current/supplier/Data_Standardized_Front_End_Data_Model.xlsx"
MR_INPUT = ROOT / "data/processed/current/market_research/Data_Standardized_MR_Front_End_Data_Model.xlsx"
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
COLUMNS = ["Dataset", "Destination", "Supplier / MR Source", "Year", "Month", "Data Type", "Period Present", "Freight Row Count", "Freight Labels", "Freight Values", "Total Freight", "Status", "Issue"]


def clean(value: Any) -> str:
    text = "" if value is None else str(value).replace("\xa0", " ").strip()
    return "" if text.casefold() in {"nan", "none", "nat"} else text


def number(value: Any) -> float | None:
    if isinstance(value, bool): return None
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else None
    text = clean(value).replace(",", "").replace("$", "")
    if text.startswith("(") and text.endswith(")"): text = f"-{text[1:-1]}"
    try:
        result = float(text)
        return result if math.isfinite(result) else None
    except ValueError:
        return None


def read_rows(path: Path) -> Iterable[dict[str, Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["front_end_data_model"] if "front_end_data_model" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(values)]
        for row in values:
            if any(value is not None and clean(value) for value in row):
                yield dict(zip(headers, row))
    finally:
        workbook.close()


def parse_period(row: dict[str, Any]) -> tuple[int, int] | None:
    try:
        year, month = int(float(clean(row.get("Time Period Year")))), int(float(clean(row.get("Time Period Month"))))
        if 1 <= month <= 12: return year, month
    except ValueError:
        pass
    value = row.get("Time_Period")
    if isinstance(value, (date, datetime)): return value.year, value.month
    for fmt in ("%Y-%m-%d", "%Y-%m", "%B %Y", "%b %Y"):
        try:
            parsed = datetime.strptime(clean(value), fmt)
            return parsed.year, parsed.month
        except ValueError:
            pass
    return None


def is_freight(row: dict[str, Any]) -> bool:
    return clean(row.get("Mapping Columns")).casefold() == "freight" or clean(row.get("Common Component")).casefold() == "freight"


def collect(dataset: str, rows: Iterable[dict[str, Any]], year_filter: int | None, entity_filter: str | None, destination_filter: str | None) -> list[dict[str, Any]]:
    periods = defaultdict(set)
    freight = defaultdict(list)
    combinations = set()
    for row in rows:
        destination, entity, period = clean(row.get("Destination Country")), clean(row.get("Supplier Name")), parse_period(row)
        if not destination or not entity or period is None: continue
        year, month = period
        if year_filter is not None and year != year_filter: continue
        if entity_filter and entity.casefold() != entity_filter.casefold(): continue
        if destination_filter and destination.casefold() != destination_filter.casefold(): continue
        key = (destination, entity, year, month)
        combinations.add((destination, entity, year))
        periods[key].add(clean(row.get("Data Type")) or "Unknown")
        if is_freight(row):
            label = clean(row.get("Raw Cost Breakdown")) or clean(row.get("Mapping Columns")) or "Freight"
            freight[key].append((label, row.get("Value")))
    output = []
    for destination, entity, year in sorted(combinations, key=lambda item: (item[2], item[0], item[1])):
        for month in range(1, 13):
            key, items = (destination, entity, year, month), freight.get((destination, entity, year, month), [])
            parsed = [number(value) for _, value in items]
            invalid = [label for (label, _), value in zip(items, parsed) if value is None]
            negative = [label for (label, _), value in zip(items, parsed) if value is not None and value < 0]
            present = key in periods
            if not present: status, issue = "FAIL", "Month is missing from the dataset"
            elif not items: status, issue = "FAIL", "No freight row for this month"
            elif invalid: status, issue = "FAIL", f"Non-numeric freight: {', '.join(invalid)}"
            elif negative: status, issue = "FAIL", f"Negative freight: {', '.join(negative)}"
            else: status, issue = "PASS", ""
            valid = [value for value in parsed if value is not None]
            output.append({"Dataset": dataset, "Destination": destination, "Supplier / MR Source": entity, "Year": year, "Month": MONTHS[month-1], "Data Type": ", ".join(sorted(periods.get(key, set()))), "Period Present": "Yes" if present else "No", "Freight Row Count": len(items), "Freight Labels": " | ".join(label for label, _ in items), "Freight Values": " | ".join(clean(value) for _, value in items), "Total Freight": sum(valid) if valid and not invalid else None, "Status": status, "Issue": issue})
    return output


def write_excel(output: Path, rows: list[dict[str, Any]], supplier_input: Path, mr_input: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(); summary = workbook.active; summary.title = "Summary"
    blue, green, red = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    summary.append(["Dataset", "Input File", "Months Checked", "Passed", "Failed"])
    for dataset, source in (("Supplier", supplier_input), ("Market Research", mr_input)):
        selected = [row for row in rows if row["Dataset"] == dataset]
        summary.append([dataset, str(source), len(selected), sum(row["Status"] == "PASS" for row in selected), sum(row["Status"] == "FAIL" for row in selected)])
    summary.append(["TOTAL", "", len(rows), sum(row["Status"] == "PASS" for row in rows), sum(row["Status"] == "FAIL" for row in rows)])
    def add_sheet(name: str, selected: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(name); sheet.append(COLUMNS)
        for row in selected:
            sheet.append([row[column] for column in COLUMNS])
            status = sheet.cell(sheet.max_row, COLUMNS.index("Status") + 1); status.fill = green if status.value == "PASS" else red; status.font = Font(bold=True)
            sheet.cell(sheet.max_row, COLUMNS.index("Total Freight") + 1).number_format = '#,##0.00'
        sheet.freeze_panes = "G2"; sheet.auto_filter.ref = sheet.dimensions
        for index, width in enumerate([18,24,38,10,13,14,16,18,55,34,17,12,48], 1): sheet.column_dimensions[get_column_letter(index)].width = width
    add_sheet("Supplier Freight", [r for r in rows if r["Dataset"] == "Supplier"])
    add_sheet("Market Freight", [r for r in rows if r["Dataset"] == "Market Research"])
    add_sheet("Failures", [r for r in rows if r["Status"] == "FAIL"])
    for sheet in workbook.worksheets:
        for cell in sheet[1]: cell.fill = blue; cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = sheet.freeze_panes or "A2"; sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate([20,95,18,14,14], 1): summary.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate monthly freight values in Supplier and Market Research standardized models.")
    parser.add_argument("--supplier-input", type=Path, default=SUPPLIER_INPUT); parser.add_argument("--mr-input", type=Path, default=MR_INPUT)
    parser.add_argument("--year", type=int); parser.add_argument("--supplier", help="Supplier Name or MR source country")
    parser.add_argument("--destination", "--market", dest="destination"); parser.add_argument("--output-excel", type=Path)
    args = parser.parse_args(); output = args.output_excel or Path(f"monthly_freight_validation_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    for label, path in (("Supplier", args.supplier_input), ("Market Research", args.mr_input)):
        if not path.is_file(): print(f"ERROR: {label} input not found: {path}", file=sys.stderr); return 2
    try:
        rows = collect("Supplier", read_rows(args.supplier_input), args.year, args.supplier, args.destination)
        rows += collect("Market Research", read_rows(args.mr_input), args.year, args.supplier, args.destination)
        if not rows: print("ERROR: No combinations matched the filters.", file=sys.stderr); return 2
        write_excel(output, rows, args.supplier_input, args.mr_input)
    except (OSError, ValueError, ImportError, StopIteration) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); return 2
    failures = sum(row["Status"] == "FAIL" for row in rows)
    print(f"Months checked: {len(rows)} | Passed: {len(rows)-failures} | Failed: {failures}")
    print(f"Excel report: {output.resolve()}")
    return 1 if failures else 0


if __name__ == "__main__": raise SystemExit(main())