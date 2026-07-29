"""Validate that displayed component differences add to the final TLC difference.

Run: python testing_scripts/validate_component_difference_totals.py --year 2026
"""
from __future__ import annotations
import argparse, sys
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "apps" / "backend"
TEST_DIR = Path(__file__).resolve().parent
MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")
COMPONENTS = ("Resin Index", "Freight", "Insurance", "Duty & Import Taxes", "Local Taxes & Fees", "Logistics & Other Costs")
COLUMNS = ["Destination", "Market Research Source", "Supplier", "Supplier Location", "Year", "Month", "Market Data Type", "Supplier Data Type", *[f"{x} Difference" for x in COMPONENTS], "Sum of Component Differences", "Market TLC", "Supplier TLC", "Final TLC Difference", "Variance", "Variance %", "Tolerance %", "Status", "Calculation"]


def total_from_rows(rows: list[dict[str, Any]], mapper) -> float | None:
    from validate_tlc_source_view_formula import number
    values = [number(row.get("amount")) for row in rows if mapper(row) == "Total Landed Cost (PET Resin)"]
    valid = [value for value in values if value is not None]
    return sum(valid) if valid else None


def supplier_values(entry: dict[str, Any], mapper) -> tuple[dict[str, float], float | None]:
    from validate_tlc_source_view_formula import bolivia_formula, component_values, cristalpet_formula, source_tlc, uruguay_cristalpet_formula, valgroup_formula
    rows = entry.get("rows", [])
    supplier = str(entry.get("supplier") or entry.get("supplierName") or "").strip().casefold()
    destination = str(entry.get("destination") or "").strip().casefold()
    if supplier == "valgroup":
        _, _, components = valgroup_formula(rows)
    elif supplier == "cristalpet" and destination == "brazil":
        _, _, components = cristalpet_formula(rows)
    elif destination == "bolivia":
        _, _, components = bolivia_formula(rows)
    elif supplier == "cristalpet" and destination == "uruguay":
        _, _, components = uruguay_cristalpet_formula(rows)
    else:
        components, _ = component_values(rows, mapper)
    total, _ = source_tlc(rows)
    return components, total if total is not None else total_from_rows(rows, mapper)


def make_row(destination, country, supplier, market_mapper, supplier_mapper, tolerance):
    from validate_tlc_source_view_formula import component_values, rounded
    market_rows = country.get("breakdown", [])
    market_components, _ = component_values(market_rows, market_mapper)
    supplier_components, supplier_tlc = supplier_values(supplier, supplier_mapper)
    market_tlc = total_from_rows(market_rows, market_mapper)
    diffs = {name: rounded(market_components.get(name, 0) - supplier_components.get(name, 0), 2) for name in COMPONENTS}
    component_sum = rounded(sum(value or 0 for value in diffs.values()), 2)
    final_diff = rounded(market_tlc - supplier_tlc, 2) if market_tlc is not None and supplier_tlc is not None else None
    variance = rounded(abs(component_sum - final_diff), 2) if final_diff is not None else None
    if variance is None:
        variance_percent = None
    elif final_diff == 0:
        variance_percent = 0.0 if variance == 0 else None
    else:
        variance_percent = rounded((variance / abs(final_diff)) * 100, 4)
    return {"Destination": destination, "Market Research Source": str(country.get("country") or ""), "Supplier": str(supplier.get("supplierName") or supplier.get("supplier") or ""), "Supplier Location": str(supplier.get("location") or ""), "Year": int(supplier.get("year") or 0), "Month": str(supplier.get("month") or ""), "Market Data Type": str(country.get("dataType") or ""), "Supplier Data Type": str(supplier.get("dataType") or ""), **{f"{name} Difference": diffs[name] for name in COMPONENTS}, "Sum of Component Differences": component_sum, "Market TLC": market_tlc, "Supplier TLC": supplier_tlc, "Final TLC Difference": final_diff, "Variance": variance, "Variance %": variance_percent, "Tolerance %": tolerance, "Status": "PASS" if variance_percent is not None and variance_percent <= tolerance else "FAIL", "Calculation": " + ".join(f"{diffs[name]:.2f}" for name in COMPONENTS) + f" = {component_sum:.2f}"}


def collect_rows(args):
    from market_research_data_model import available_destinations, available_periods_for_destination, build_market_research_countries
    from supplier_data_model import build_vendor_breakdowns
    from validate_tlc_breakdown import map_market_row, map_supplier_row
    entries = build_vendor_breakdowns(("China",))
    destinations = available_destinations()
    if args.destination:
        destinations = [d for d in destinations if d.casefold() == args.destination.casefold()]
    output, seen = [], set()
    for destination in destinations:
        for month, year_text in available_periods_for_destination(destination):
            year = int(year_text)
            if args.year is not None and year != args.year:
                continue
            countries = build_market_research_countries(destination, month, year)
            if args.market_source:
                countries = [c for c in countries if str(c.get("country", "")).casefold() == args.market_source.casefold()]
            suppliers = [e for e in entries if str(e.get("destination", "")).casefold() == destination.casefold() and str(e.get("month", "")).casefold() == month.casefold() and int(e.get("year") or 0) == year and (not args.supplier or str(e.get("supplierName") or e.get("supplier") or "").casefold() == args.supplier.casefold())]
            for country in countries:
                for supplier in suppliers:
                    key = (destination, country.get("country"), supplier.get("supplierName"), supplier.get("location"), year, month, supplier.get("dataType"))
                    if key not in seen:
                        seen.add(key)
                        output.append(make_row(destination, country, supplier, map_market_row, map_supplier_row, args.tolerance))
    output.sort(key=lambda r: (r["Destination"], r["Year"], MONTHS.index(r["Month"]) if r["Month"] in MONTHS else 99, r["Market Research Source"], r["Supplier"]))
    return output


def write_excel(rows, output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); summary = wb.active; summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Combinations checked", len(rows)])
    summary.append(["Passed", sum(r["Status"] == "PASS" for r in rows)])
    summary.append(["Failed", sum(r["Status"] == "FAIL" for r in rows)])
    summary.append(["Maximum variance", max((r["Variance"] or 0 for r in rows), default=0)])
    summary.append(["Maximum variance %", max((r["Variance %"] or 0 for r in rows), default=0)])
    summary.append(["Tolerance %", rows[0]["Tolerance %"] if rows else None])
    def add_sheet(name, selected):
        ws = wb.create_sheet(name); ws.append(COLUMNS)
        for row in selected:
            ws.append([row.get(c) for c in COLUMNS])
            status = ws.cell(ws.max_row, COLUMNS.index("Status") + 1)
            status.fill = PatternFill("solid", fgColor="C6EFCE" if status.value == "PASS" else "FFC7CE"); status.font = Font(bold=True)
            for col in range(9, COLUMNS.index("Tolerance %") + 2): ws.cell(ws.max_row, col).number_format = '#,##0.00;[Red]-#,##0.00'
        ws.freeze_panes = "I2"; ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(COLUMNS, 1): ws.column_dimensions[get_column_letter(i)].width = 52 if col == "Calculation" else max(12, min(25, len(col) + 2))
    add_sheet("All Combinations", rows); add_sheet("Failures", [r for r in rows if r["Status"] == "FAIL"])
    header = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        for cell in ws[1]: cell.fill = header; cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
    wb.save(output)


def main():
    parser = argparse.ArgumentParser(description="Validate sum(component differences) = final TLC difference.")
    parser.add_argument("--year", type=int); parser.add_argument("--destination"); parser.add_argument("--supplier"); parser.add_argument("--market-source")
    parser.add_argument("--tolerance", type=float, default=2.0, help="Allowed percentage variance; default 2%%")
    parser.add_argument("--output-excel", type=Path); args = parser.parse_args()
    output = args.output_excel or TEST_DIR / f"component_difference_totals_{args.year or 'all'}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    sys.path[:0] = [str(BACKEND), str(TEST_DIR)]
    try:
        rows = collect_rows(args)
        if not rows: print("ERROR: No combinations matched.", file=sys.stderr); return 2
        write_excel(rows, output)
    except (ImportError, OSError, ValueError, KeyError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr); return 2
    failures = sum(r["Status"] == "FAIL" for r in rows)
    print(f"Combinations checked: {len(rows)} | Passed: {len(rows)-failures} | Failed: {failures}")
    print(f"Excel report: {output.resolve()}")
    return 1 if failures else 0

if __name__ == "__main__": raise SystemExit(main())



